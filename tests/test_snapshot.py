from __future__ import annotations

import hashlib
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from jsonschema import Draft202012Validator

from thesis_format.contracts import canonical_json, canonical_sha256, load_schema
from thesis_format.profile import load_profile
from thesis_format.snapshot import (
    assign_node_ids,
    compare_snapshots,
    extract_snapshot,
    normalize_text,
    strip_manual_numbering,
    write_content_reports,
)


ROOT = Path(__file__).parents[1]
PROFILE = load_profile(ROOT / "rules" / "apple-university.yaml")


def _field(paragraph, instruction: str, result: str) -> None:
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.text = instruction
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    begin_run._r.extend((begin, instruction_node, separator))
    paragraph.add_run(result)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def _save(document: Document, path: Path) -> Path:
    document.save(path)
    return path


def _add_image_relationship(
    path: Path, media: bytes = b"same-image", *, occurrences: int = 1
) -> None:
    replacement = path.with_suffix(".replacement.docx")
    with ZipFile(path) as source, ZipFile(
        replacement, "w", compression=ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                marker = b"</w:p>"
                drawing = (
                    '<w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/'
                    'drawingml/2006/wordprocessingDrawing"><wp:docPr id="7" name="Picture 7" '
                    'descr="  Cafe\u0301\u00a0tree  "/><a:graphic xmlns:a="http://schemas.'
                    'openxmlformats.org/drawingml/2006/main"><a:graphicData><pic:pic '
                    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
                    '<pic:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/'
                    'officeDocument/2006/relationships" r:embed="rIdSnapshotImage"/>'
                    "</pic:blipFill></pic:pic></a:graphicData></a:graphic></wp:inline>"
                    "</w:drawing></w:r>"
                ).encode()
                data = data.replace(marker, drawing * occurrences + marker, 1)
            elif info.filename == "word/_rels/document.xml.rels":
                relationship = (
                    b'<Relationship Id="rIdSnapshotImage" '
                    b'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    b'relationships/image" Target="media/snapshot.bin"/>'
                )
                data = data.replace(b"</Relationships>", relationship + b"</Relationships>")
            target.writestr(info, data)
        target.writestr("word/media/snapshot.bin", media)
    replacement.replace(path)


def _append_equation(paragraph, text: str) -> None:
    math_paragraph = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    math.append(run)
    math_paragraph.append(math)
    paragraph._p.append(math_paragraph)


def _append_note_reference(paragraph, kind: str, note_id: int) -> None:
    run = paragraph.add_run()
    reference = OxmlElement(f"w:{kind}Reference")
    reference.set(qn("w:id"), str(note_id))
    run._r.append(reference)


def _add_notes_part(path: Path, kind: str, notes: dict[int, str]) -> None:
    replacement = path.with_suffix(f".{kind}-replacement.docx")
    entries = "".join(
        f'<w:{kind} w:id="{note_id}"><w:p><w:r><w:t>{text}</w:t>'
        f"</w:r></w:p></w:{kind}>"
        for note_id, text in notes.items()
    )
    content = (
        f'<w:{kind}s xmlns:w="http://schemas.openxmlformats.org/'
        f'wordprocessingml/2006/main">{entries}</w:{kind}s>'
    ).encode()
    with ZipFile(path) as source, ZipFile(
        replacement, "w", compression=ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        target.writestr(f"word/{kind}s.xml", content)
    replacement.replace(path)


def _png() -> BytesIO:
    return BytesIO(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
            "0000000d49444154789c6360f8cfc000000301010018dd8db10000000049454e44"
            "ae426082"
        )
    )


def _snapshot(nodes: list[dict[str, object]]) -> dict[str, object]:
    return _snapshot_stories(body=nodes)


def _snapshot_stories(
    *,
    body: list[dict[str, object]] | None = None,
    headers: list[dict[str, object]] | None = None,
    footers: list[dict[str, object]] | None = None,
) -> dict[str, object]:
    return {
        "schema_version": "1.0",
        "snapshot_id": "0" * 16,
        "document_sha256": "a" * 64,
        "body": body or [],
        "headers": headers or [],
        "footers": footers or [],
    }


def _node(
    node_id: str,
    text: str,
    index: int,
    *,
    node_type: str = "paragraph",
    relationship_hash: str | None = None,
) -> dict[str, object]:
    return {
        "node_id": node_id,
        "node_type": node_type,
        "source_location": {
            "part": "word/document.xml",
            "path": f"/w:document/w:body/*[{index + 1}]",
            "index": index,
        },
        "normalized_text": text,
        "relationship_hash": relationship_hash,
    }


def test_normalizes_unicode_line_endings_and_all_whitespace() -> None:
    assert normalize_text("  Cafe\u0301\r\n\u00a0\t中文  ") == "Café 中文"
    assert normalize_text("，\u2003。") == "， 。"


def test_strips_manual_numbering_only_for_confirmed_profile_role() -> None:
    assert strip_manual_numbering("1.2 研究背景", "heading_2", PROFILE) == "研究背景"
    assert strip_manual_numbering("1.2 研究背景", None, PROFILE) == "1.2 研究背景"
    assert strip_manual_numbering("图 1-2 苹果树", "figure_caption", PROFILE) == "苹果树"
    assert strip_manual_numbering("[12] Apple study.", "reference", PROFILE) == "Apple study."


def test_extracts_fields_tables_references_headers_footers_and_empty_cells(
    tmp_path: Path,
) -> None:
    document = Document()
    document.add_paragraph("")
    toc = document.add_paragraph("目录 ")
    _field(toc, ' TOC \\o "1-3" ', "第一章 绪论 1")
    caption = document.add_paragraph("图 ")
    _field(caption, " SEQ Figure ", "1-1")
    caption.add_run(" 苹果树")
    prose = document.add_paragraph("结果见")
    _field(prose, " REF target ", "图 1-1")
    prose.add_run("。")
    document.add_paragraph("参考文献")
    reference = document.add_paragraph("[1] Apple Research Group.")
    reference.style = document.styles["Normal"]
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = ""
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"
    document.sections[0].header.paragraphs[0].text = "Header One"
    footer = document.sections[0].footer.paragraphs[0]
    _field(footer, " PAGE ", "9")
    footer.add_run(" Footer")
    path = _save(document, tmp_path / "story.docx")

    snapshot = extract_snapshot(path, PROFILE)

    Draft202012Validator(load_schema("content-snapshot")).validate(snapshot)
    texts = [node["normalized_text"] for node in snapshot["body"]]
    assert "" not in [
        node["normalized_text"]
        for node in snapshot["body"]
        if node["node_type"] == "paragraph"
    ]
    assert texts[:5] == ["目录", "苹果树", "结果见。", "参考文献", "Apple Research Group."]
    assert snapshot["body"][1]["node_type"] == "caption"
    cells = [node["normalized_text"] for node in snapshot["body"] if node["node_type"] == "table"]
    assert cells == ["A", "", "C", "D"]
    assert [node["normalized_text"] for node in snapshot["headers"]] == ["Header One"]
    assert [node["normalized_text"] for node in snapshot["footers"]] == ["Footer"]


def test_extracts_figure_alt_text_and_target_sha256(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("anchor")
    path = _save(document, tmp_path / "figure.docx")
    _add_image_relationship(path)

    snapshot = extract_snapshot(path, PROFILE)
    image = next(node for node in snapshot["body"] if node["node_type"] == "image")

    assert image["normalized_text"] == "Café tree"
    content_hash = hashlib.sha256(b"same-image").hexdigest()
    assert image["relationship_hash"] == canonical_sha256(
        [content_hash, "body", 1]
    )


def test_repeated_identical_images_and_tables_get_distinct_occurrence_identities(
    tmp_path: Path,
) -> None:
    document = Document()
    document.add_paragraph("anchor")
    for _ in range(2):
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "same cell"
    path = _save(document, tmp_path / "repeated-objects.docx")
    _add_image_relationship(path, occurrences=2)

    snapshot = extract_snapshot(path, PROFILE)
    images = [node for node in snapshot["body"] if node["node_type"] == "image"]
    tables = [node for node in snapshot["body"] if node["node_type"] == "table"]

    assert len(images) == 2
    image_content_hash = hashlib.sha256(b"same-image").hexdigest()
    assert [node["relationship_hash"] for node in images] == [
        canonical_sha256([image_content_hash, "body", 1]),
        canonical_sha256([image_content_hash, "body", 2]),
    ]
    assert len(tables) == 2
    table_content_hash = canonical_sha256(
        {
            "row_count": 1,
            "column_count": 1,
            "cells": [
                {
                    "row": 1,
                    "column": 1,
                    "grid_span": 1,
                    "v_merge": None,
                    "text": "same cell",
                }
            ],
        }
    )
    assert [node["relationship_hash"] for node in tables] == [
        canonical_sha256([table_content_hash, "body", 1]),
        canonical_sha256([table_content_hash, "body", 2]),
    ]


def test_note_references_are_ordered_semantic_nodes_linked_to_note_content(
    tmp_path: Path,
) -> None:
    before_document = Document()
    first = before_document.add_paragraph("First")
    _append_note_reference(first, "footnote", 4)
    second = before_document.add_paragraph("Second")
    _append_note_reference(second, "endnote", 7)
    before_path = _save(before_document, tmp_path / "note-references-before.docx")
    _add_notes_part(before_path, "footnote", {4: "Footnote content"})
    _add_notes_part(before_path, "endnote", {7: "Endnote content"})

    before = extract_snapshot(before_path, PROFILE)
    assert [
        (node["node_type"], node.get("note_id"), node["normalized_text"])
        for node in before["body"]
    ] == [
        ("paragraph", None, "First"),
        ("footnote_reference", "4", ""),
        ("paragraph", None, "Second"),
        ("endnote_reference", "7", ""),
        ("footnote", "4", "Footnote content"),
        ("endnote", "7", "Endnote content"),
    ]
    assert before["body"][1]["source_location"]["path"].endswith(
        "/w:p[1]/w:footnoteReference[1]"
    )
    assert before["body"][3]["source_location"]["path"].endswith(
        "/w:p[2]/w:endnoteReference[1]"
    )

    after_document = Document()
    after_document.add_paragraph("First")
    moved = after_document.add_paragraph("Second")
    _append_note_reference(moved, "footnote", 4)
    _append_note_reference(moved, "endnote", 7)
    after_path = _save(after_document, tmp_path / "note-references-after.docx")
    _add_notes_part(after_path, "footnote", {4: "Footnote content"})
    _add_notes_part(after_path, "endnote", {7: "Endnote content"})

    comparison = compare_snapshots(before, extract_snapshot(after_path, PROFILE))

    assert comparison["equivalent"] is False
    assert any(
        difference["before"]["node_type"] == "footnote_reference"
        for difference in comparison["differences"]
        if difference["before"] is not None
    )


def test_table_identity_includes_shape_even_when_row_major_text_matches(
    tmp_path: Path,
) -> None:
    before_document = Document()
    before_table = before_document.add_table(rows=2, cols=2)
    after_document = Document()
    after_table = after_document.add_table(rows=1, cols=4)
    for table in (before_table, after_table):
        for cell, text in zip(
            (cell for row in table.rows for cell in row.cells),
            ("A", "B", "C", "D"),
            strict=True,
        ):
            cell.text = text

    before = extract_snapshot(
        _save(before_document, tmp_path / "table-2x2.docx"), PROFILE
    )
    after = extract_snapshot(
        _save(after_document, tmp_path / "table-1x4.docx"), PROFILE
    )
    before_cells = [node for node in before["body"] if node["node_type"] == "table"]
    after_cells = [node for node in after["body"] if node["node_type"] == "table"]

    assert before_cells[0]["table_geometry"]["row_count"] == 2
    assert before_cells[0]["table_geometry"]["column_count"] == 2
    assert after_cells[0]["table_geometry"]["row_count"] == 1
    assert after_cells[0]["table_geometry"]["column_count"] == 4
    assert before_cells[0]["relationship_hash"] != after_cells[0]["relationship_hash"]
    assert compare_snapshots(before, after)["equivalent"] is False


def test_table_geometry_records_spans_merges_and_logical_coordinates(
    tmp_path: Path,
) -> None:
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Horizontal"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 2).text = "Vertical"
    table.cell(1, 2).merge(table.cell(2, 2))

    snapshot = extract_snapshot(
        _save(document, tmp_path / "table-spans-merges.docx"), PROFILE
    )
    cells = [node for node in snapshot["body"] if node["node_type"] == "table"]
    geometries = [cell["table_geometry"] for cell in cells]

    assert all(geometry["row_count"] == 3 for geometry in geometries)
    assert all(geometry["column_count"] == 3 for geometry in geometries)
    assert geometries[0] == {
        "row_count": 3,
        "column_count": 3,
        "row": 1,
        "column": 1,
        "grid_span": 2,
        "v_merge": None,
    }
    assert any(
        geometry["row"] == 2
        and geometry["column"] == 3
        and geometry["v_merge"] == "restart"
        for geometry in geometries
    )
    assert any(
        geometry["row"] == 3
        and geometry["column"] == 3
        and geometry["v_merge"] == "continue"
        for geometry in geometries
    )


def test_duplicate_nodes_get_ordinals_and_prefix_collisions_extend() -> None:
    payload = ["1.0", "paragraph", "same", "", 0]
    duplicate = ["1.0", "paragraph", "same", "", 1]
    digests = iter(
        [
            "0123456789abcdef1111" + "0" * 44,
            "0123456789abcdef2222" + "0" * 44,
        ]
    )

    ids = assign_node_ids([payload, duplicate], digest_function=lambda _: next(digests))

    assert ids == ["0123456789abcdef1111", "0123456789abcdef2222"]


def test_duplicate_paragraphs_and_section_stories_remain_distinct(
    tmp_path: Path,
) -> None:
    document = Document()
    document.add_paragraph("Repeated")
    document.add_paragraph("Repeated")
    first = document.sections[0]
    first.header.paragraphs[0].text = "First header"
    first.footer.paragraphs[0].text = "First footer"
    second = document.add_section()
    second.header.is_linked_to_previous = False
    second.footer.is_linked_to_previous = False
    second.header.paragraphs[0].text = "Second header"
    second.footer.paragraphs[0].text = "Second footer"

    snapshot = extract_snapshot(_save(document, tmp_path / "sections.docx"), PROFILE)

    repeated_ids = [
        node["node_id"]
        for node in snapshot["body"]
        if node["normalized_text"] == "Repeated"
    ]
    assert len(repeated_ids) == 2
    assert len(set(repeated_ids)) == 2
    assert [node["normalized_text"] for node in snapshot["headers"]] == [
        "First header",
        "Second header",
    ]
    assert [node["normalized_text"] for node in snapshot["footers"]] == [
        "First footer",
        "Second footer",
    ]
    assert snapshot["headers"][0]["source_location"]["path"].startswith("/section[1]/")
    assert snapshot["headers"][1]["source_location"]["path"].startswith("/section[2]/")


def test_compare_allows_only_declared_caption_move_adjacent_to_same_object() -> None:
    object_hash = "b" * 64
    caption = _node("c" * 16, "Apple tree", 0, node_type="caption")
    image = _node(
        "d" * 16,
        "Apple tree image",
        1,
        node_type="image",
        relationship_hash=object_hash,
    )
    before = _snapshot([caption, image])
    after = _snapshot(
        [
            {**image, "source_location": {**image["source_location"], "index": 0}},
            {**caption, "source_location": {**caption["source_location"], "index": 1}},
        ]
    )

    comparison = compare_snapshots(
        before,
        after,
        changes=[
            {
                "rule_id": "FIGURE.CAPTION.POSITION",
                "operation": "move_caption",
                "status": "applied",
                "node_id": caption["node_id"],
                "relationship_sha256": object_hash,
            }
        ],
    )

    assert comparison["equivalent"] is True
    assert comparison["differences"] == []


@pytest.mark.parametrize(
    ("object_type", "rule_id"),
    [
        ("image", "FIGURE.CAPTION.POSITION"),
        ("table", "TABLE.CAPTION.POSITION"),
    ],
)
def test_compare_rejects_caption_move_to_wrong_identical_object_occurrence(
    object_type: str,
    rule_id: str,
) -> None:
    first_identity = "1" * 64
    second_identity = "2" * 64
    first_caption = _node("a" * 16, "First caption", 0, node_type="caption")
    first_object = _node(
        "b" * 16,
        "same object",
        1,
        node_type=object_type,
        relationship_hash=first_identity,
    )
    second_caption = _node("c" * 16, "Second caption", 2, node_type="caption")
    second_object = _node(
        "d" * 16,
        "same object",
        3,
        node_type=object_type,
        relationship_hash=second_identity,
    )
    after_nodes = [
        first_object,
        second_caption,
        second_object,
        first_caption,
    ]
    after = [
        {**node, "source_location": {**node["source_location"], "index": index}}
        for index, node in enumerate(after_nodes)
    ]

    comparison = compare_snapshots(
        _snapshot([first_caption, first_object, second_caption, second_object]),
        _snapshot(after),
        changes=[
            {
                "rule_id": rule_id,
                "operation": "move_caption",
                "status": "applied",
                "node_id": first_caption["node_id"],
                "relationship_sha256": first_identity,
            }
        ],
    )

    assert comparison["equivalent"] is False
    assert {difference["kind"] for difference in comparison["differences"]} == {
        "reordered"
    }


@pytest.mark.parametrize(
    "change",
    [
        {
            "rule_id": "FIGURE.CAPTION.POSITION",
            "operation": "move_caption",
            "status": "applied",
            "relationship_sha256": "b" * 64,
        },
        {
            "rule_id": "FIGURE.CAPTION.POSITION",
            "operation": "move_caption",
            "status": "applied",
            "node_id": "e" * 16,
            "relationship_sha256": "b" * 64,
        },
        {
            "rule_id": "FIGURE.CAPTION.POSITION",
            "operation": "move_caption",
            "status": "applied",
            "node_id": "c" * 16,
            "relationship_sha256": "f" * 64,
        },
        {
            "rule_id": "FIGURE.UNRELATED",
            "operation": "move_caption",
            "status": "applied",
            "node_id": "c" * 16,
            "relationship_sha256": "b" * 64,
        },
    ],
)
def test_compare_rejects_caption_move_without_exact_transformation_binding(
    change: dict[str, object],
) -> None:
    object_hash = "b" * 64
    caption = _node("c" * 16, "Apple tree", 0, node_type="caption")
    image = _node(
        "d" * 16,
        "Apple tree image",
        1,
        node_type="image",
        relationship_hash=object_hash,
    )
    after = [
        {**image, "source_location": {**image["source_location"], "index": 0}},
        {**caption, "source_location": {**caption["source_location"], "index": 1}},
    ]

    comparison = compare_snapshots(
        _snapshot([caption, image]),
        _snapshot(after),
        changes=[change],
    )

    assert comparison["equivalent"] is False
    assert {difference["kind"] for difference in comparison["differences"]} == {
        "reordered"
    }


def test_compare_rejects_caption_exemption_for_paragraph_or_nonadjacent_object() -> None:
    object_hash = "b" * 64
    paragraph = _node("c" * 16, "Apple tree", 0)
    image = _node(
        "d" * 16,
        "Apple tree image",
        1,
        node_type="image",
        relationship_hash=object_hash,
    )
    spacer = _node("e" * 16, "Unrelated prose", 2)
    change = {
        "rule_id": "FIGURE.CAPTION.POSITION",
        "operation": "move_caption",
        "status": "applied",
        "node_id": paragraph["node_id"],
        "relationship_sha256": object_hash,
    }

    paragraph_comparison = compare_snapshots(
        _snapshot([paragraph, image]),
        _snapshot(
            [
                {**image, "source_location": {**image["source_location"], "index": 0}},
                {**paragraph, "source_location": {**paragraph["source_location"], "index": 1}},
            ]
        ),
        changes=[change],
    )
    caption = {**paragraph, "node_type": "caption"}
    nonadjacent_comparison = compare_snapshots(
        _snapshot([caption, spacer, image]),
        _snapshot(
            [
                {**image, "source_location": {**image["source_location"], "index": 0}},
                spacer,
                {**caption, "source_location": {**caption["source_location"], "index": 2}},
            ]
        ),
        changes=[change],
    )

    assert paragraph_comparison["equivalent"] is False
    assert nonadjacent_comparison["equivalent"] is False


def test_extracts_notes_and_equations_deterministically(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Body")
    path = _save(document, tmp_path / "notes-and-equations.docx")
    replacement = path.with_suffix(".replacement.docx")
    math = (
        b'<w:p><m:oMathPara xmlns:m="http://schemas.openxmlformats.org/'
        b'officeDocument/2006/math"><m:oMath><m:r><m:t>x + 1</m:t></m:r>'
        b"</m:oMath></m:oMathPara></w:p>"
    )
    footnotes = (
        b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/'
        b'wordprocessingml/2006/main"><w:footnote w:id="-1"><w:p><w:r>'
        b"<w:t>separator</w:t></w:r></w:p></w:footnote>"
        b'<w:footnote w:id="2"><w:p><w:r><w:t>Second note</w:t></w:r></w:p>'
        b'</w:footnote><w:footnote w:id="1"><w:p><w:r><w:t>First note</w:t>'
        b"</w:r></w:p></w:footnote></w:footnotes>"
    )
    endnotes = (
        b'<w:endnotes xmlns:w="http://schemas.openxmlformats.org/'
        b'wordprocessingml/2006/main"><w:endnote w:id="3"><w:p><w:r>'
        b"<w:t>End note</w:t></w:r></w:p></w:endnote></w:endnotes>"
    )
    with ZipFile(path) as source, ZipFile(
        replacement, "w", compression=ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                data = data.replace(b"<w:sectPr", math + b"<w:sectPr", 1)
            target.writestr(info, data)
        target.writestr("word/footnotes.xml", footnotes)
        target.writestr("word/endnotes.xml", endnotes)
    replacement.replace(path)

    first = extract_snapshot(path, PROFILE)
    second = extract_snapshot(path, PROFILE)

    assert first == second
    assert [
        (node["node_type"], node["normalized_text"])
        for node in first["body"]
    ] == [
        ("paragraph", "Body"),
        ("equation", "x + 1"),
        ("footnote", "First note"),
        ("footnote", "Second note"),
        ("endnote", "End note"),
    ]


def test_resolves_each_header_footer_type_independently_across_sections(
    tmp_path: Path,
) -> None:
    document = Document()
    document.settings.odd_and_even_pages_header_footer = True
    first = document.sections[0]
    first.different_first_page_header_footer = True
    first.header.paragraphs[0].text = "S1 default header"
    first.first_page_header.paragraphs[0].text = "S1 first header"
    first.even_page_header.paragraphs[0].text = "S1 even header"
    first.footer.paragraphs[0].text = "S1 default footer"
    first.first_page_footer.paragraphs[0].text = "S1 first footer"
    first.even_page_footer.paragraphs[0].text = "S1 even footer"

    second = document.add_section(WD_SECTION_START.NEW_PAGE)
    second.different_first_page_header_footer = True
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = "S2 default header"
    second.first_page_footer.is_linked_to_previous = False
    second.first_page_footer.paragraphs[0].text = "S2 first footer"

    snapshot = extract_snapshot(
        _save(document, tmp_path / "typed-inheritance.docx"), PROFILE
    )

    assert [node["normalized_text"] for node in snapshot["headers"]] == [
        "S1 default header",
        "S1 first header",
        "S1 even header",
        "S2 default header",
        "S1 first header",
        "S1 even header",
    ]
    assert [node["normalized_text"] for node in snapshot["footers"]] == [
        "S1 default footer",
        "S1 first footer",
        "S1 even footer",
        "S1 default footer",
        "S2 first footer",
        "S1 even footer",
    ]
    assert "/header[default]/" in snapshot["headers"][3]["source_location"]["path"]
    assert "/header[first]/" in snapshot["headers"][4]["source_location"]["path"]
    assert "/footer[first]/" in snapshot["footers"][4]["source_location"]["path"]


def test_extracts_typed_story_text_images_equations_and_tables(
    tmp_path: Path,
) -> None:
    document = Document()
    header = document.sections[0].first_page_header
    document.sections[0].different_first_page_header_footer = True
    header.paragraphs[0].text = "Header text"
    image_run = header.add_paragraph().add_run()
    inline = image_run.add_picture(_png(), width=Inches(0.1))
    inline._inline.docPr.set("descr", "Header image")
    _append_equation(header.add_paragraph(), "h + 1")
    header_table = header.add_table(rows=1, cols=1, width=Inches(1))
    header_table.cell(0, 0).text = "Header cell"

    footer = document.sections[0].even_page_footer
    document.settings.odd_and_even_pages_header_footer = True
    footer.paragraphs[0].text = "Footer text"
    _append_equation(footer.add_paragraph(), "f + 1")
    footer_table = footer.add_table(rows=1, cols=1, width=Inches(1))
    footer_table.cell(0, 0).text = "Footer cell"

    snapshot = extract_snapshot(
        _save(document, tmp_path / "typed-story-content.docx"), PROFILE
    )

    assert [
        (node["node_type"], node["normalized_text"])
        for node in snapshot["headers"]
    ] == [
        ("header", "Header text"),
        ("image", "Header image"),
        ("equation", "h + 1"),
        ("table", "Header cell"),
    ]
    assert [
        (node["node_type"], node["normalized_text"])
        for node in snapshot["footers"]
    ] == [
        ("footer", "Footer text"),
        ("equation", "f + 1"),
        ("table", "Footer cell"),
    ]


def test_compare_detects_header_image_and_footer_equation_changes(
    tmp_path: Path,
) -> None:
    before_document = Document()
    header_run = before_document.sections[0].header.paragraphs[0].add_run()
    inline = header_run.add_picture(_png(), width=Inches(0.1))
    inline._inline.docPr.set("descr", "Header image")
    _append_equation(before_document.sections[0].footer.paragraphs[0], "x + 1")
    before_path = _save(before_document, tmp_path / "typed-before.docx")

    after_path = tmp_path / "typed-after.docx"
    after_path.write_bytes(before_path.read_bytes())
    replacement = after_path.with_suffix(".replacement.docx")
    with ZipFile(after_path) as source, ZipFile(
        replacement, "w", compression=ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("word/media/"):
                data += b"changed"
            if info.filename.startswith("word/footer") and info.filename.endswith(".xml"):
                data = data.replace(b"x + 1", b"x + 2")
            target.writestr(info, data)
    replacement.replace(after_path)

    comparison = compare_snapshots(
        extract_snapshot(before_path, PROFILE),
        extract_snapshot(after_path, PROFILE),
    )

    assert comparison["equivalent"] is False
    changed_types = {
        difference["before"]["node_type"]
        for difference in comparison["differences"]
        if difference["kind"] == "changed"
    }
    assert changed_types == {"image", "equation"}


def test_compare_treats_body_to_header_move_as_remove_and_add() -> None:
    body_node = _node("1" * 16, "Moved text", 0)
    header_node = {
        **body_node,
        "node_type": "header",
        "source_location": {
            "part": "word/header1.xml",
            "path": "/w:hdr/w:p[1]",
            "index": 0,
        },
    }

    comparison = compare_snapshots(
        _snapshot_stories(body=[body_node]),
        _snapshot_stories(headers=[header_node]),
    )

    assert comparison["equivalent"] is False
    assert [difference["kind"] for difference in comparison["differences"]] == [
        "removed",
        "added",
    ]


def test_compare_rejects_prose_change_and_reports_explicit_kinds() -> None:
    before = _snapshot([_node("1" * 16, "Original prose.", 0)])
    after = _snapshot([_node("2" * 16, "Changed prose.", 0)])

    comparison = compare_snapshots(before, after)

    assert comparison["equivalent"] is False
    assert comparison["differences"][0]["kind"] == "changed"
    assert comparison["differences"][0]["before"]["normalized_text"] == "Original prose."
    assert comparison["differences"][0]["after"]["normalized_text"] == "Changed prose."


def test_compare_reports_added_removed_and_reordered() -> None:
    first = _node("1" * 16, "First", 0)
    second = _node("2" * 16, "Second", 1)
    removed = _node("3" * 16, "Removed", 2)
    added = _node("4" * 16, "Added", 2)
    before = _snapshot([first, second, removed])
    after = _snapshot(
        [
            {**second, "source_location": {**second["source_location"], "index": 0}},
            {**first, "source_location": {**first["source_location"], "index": 1}},
            added,
        ]
    )

    kinds = {difference["kind"] for difference in compare_snapshots(before, after)["differences"]}

    assert kinds == {"added", "removed", "reordered"}


def test_writes_canonical_snapshots_and_markdown_diff(tmp_path: Path) -> None:
    before = _snapshot([_node("1" * 16, "Before", 0)])
    after = _snapshot([_node("2" * 16, "After", 0)])
    comparison = compare_snapshots(before, after)

    paths = write_content_reports(tmp_path, before, after, comparison)

    assert paths == {
        "before": tmp_path / "content-before.json",
        "after": tmp_path / "content-after.json",
        "diff": tmp_path / "content-diff.md",
    }
    assert paths["before"].read_text(encoding="utf-8") == canonical_json(before)
    assert json.loads(paths["after"].read_text(encoding="utf-8")) == after
    markdown = paths["diff"].read_text(encoding="utf-8")
    assert "# Content preservation diff" in markdown
    assert "changed" in markdown
