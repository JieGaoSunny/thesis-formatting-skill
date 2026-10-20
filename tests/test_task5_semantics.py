from __future__ import annotations

from copy import deepcopy
from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from thesis_format.audit import (
    _finding_digest,
    audit_document,
    render_audit_markdown,
)
from thesis_format.contracts import validate_audit
from thesis_format.profile import load_profile
from thesis_format.semantic import inspect_document

from helpers.docx_factory import build_docx


ROOT = Path(__file__).parents[1]
PROFILE = load_profile(ROOT / "rules" / "apple-university.yaml")
NOW = datetime(2026, 10, 20, 10, 0, tzinfo=UTC)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS = {"w": W_NS, "r": R_NS, "wp": WP_NS}
W = f"{{{W_NS}}}"
R = f"{{{R_NS}}}"


def _replace_part(path: Path, name: str, root: etree._Element) -> None:
    replacement = path.with_suffix(".replacement.docx")
    data = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with ZipFile(path) as source, ZipFile(replacement, "w") as target:
        for info in source.infolist():
            target.writestr(
                info,
                data if info.filename == name else source.read(info.filename),
            )
    replacement.replace(path)


def _xml(path: Path, name: str) -> etree._Element:
    with ZipFile(path) as archive:
        return etree.fromstring(archive.read(name))


def _rule_ids(path: Path) -> list[str]:
    report = audit_document(path, PROFILE, clock=lambda: NOW)
    return [finding["rule_id"] for finding in report["findings"]]


def _set_num_format(path: Path, num_id: str, value: str) -> None:
    root = _xml(path, "word/numbering.xml")
    abstract_id = root.xpath(
        f"./w:num[@w:numId='{num_id}']/w:abstractNumId/@w:val",
        namespaces=NS,
    )[0]
    for level in root.xpath(
        f"./w:abstractNum[@w:abstractNumId='{abstract_id}']/w:lvl",
        namespaces=NS,
    ):
        num_fmt = level.find("w:numFmt", namespaces=NS)
        if num_fmt is None:
            num_fmt = etree.SubElement(level, f"{W}numFmt")
        num_fmt.set(f"{W}val", value)
    _replace_part(path, "word/numbering.xml", root)


def _set_section_page_number(
    path: Path, section: int, *, start: int | None, fmt: str | None
) -> None:
    root = _xml(path, "word/document.xml")
    sections = root.xpath(".//w:sectPr", namespaces=NS)
    pg_num = sections[section - 1].find("w:pgNumType", namespaces=NS)
    if pg_num is None:
        pg_num = etree.SubElement(sections[section - 1], f"{W}pgNumType")
    for name, value in (("start", start), ("fmt", fmt)):
        if value is None:
            pg_num.attrib.pop(f"{W}{name}", None)
        else:
            pg_num.set(f"{W}{name}", str(value))
    _replace_part(path, "word/document.xml", root)


def _add_drawing(path: Path, paragraph: int) -> None:
    root = _xml(path, "word/document.xml")
    target = root.xpath(".//w:body/w:p", namespaces=NS)[paragraph - 1]
    run = etree.SubElement(target, f"{W}r")
    drawing = etree.SubElement(run, f"{W}drawing")
    inline = etree.SubElement(
        drawing,
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline",
    )
    etree.SubElement(
        inline,
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent",
        cx="6480000",
        cy="1800000",
    )
    _replace_part(path, "word/document.xml", root)


def _set_paragraph_spacing(
    path: Path, paragraph: int, *, line: int, line_rule: str
) -> None:
    root = _xml(path, "word/document.xml")
    target = root.xpath(".//w:body/w:p", namespaces=NS)[paragraph - 1]
    p_pr = target.find("w:pPr", namespaces=NS)
    if p_pr is None:
        p_pr = etree.Element(f"{W}pPr")
        target.insert(0, p_pr)
    spacing = p_pr.find("w:spacing", namespaces=NS)
    if spacing is None:
        spacing = etree.SubElement(p_pr, f"{W}spacing")
    spacing.set(f"{W}line", str(line))
    spacing.set(f"{W}lineRule", line_rule)
    spacing.set(f"{W}before", "0")
    spacing.set(f"{W}after", "0")
    ind = p_pr.find("w:ind", namespaces=NS)
    if ind is None:
        ind = etree.SubElement(p_pr, f"{W}ind")
    ind.set(f"{W}hanging", "567")
    _replace_part(path, "word/document.xml", root)


def _set_style_to_profile(path: Path, style_id: str, role: str) -> None:
    root = _xml(path, "word/styles.xml")
    style = root.xpath(
        f"./w:style[@w:styleId='{style_id}']",
        namespaces=NS,
    )[0]
    for name in ("pPr", "rPr"):
        existing = style.find(f"w:{name}", namespaces=NS)
        if existing is not None:
            style.remove(existing)

    expected = PROFILE.styles[role]
    paragraph = expected.paragraph
    p_pr = etree.SubElement(style, f"{W}pPr")
    etree.SubElement(p_pr, f"{W}jc", {f"{W}val": paragraph.alignment})
    spacing_attributes = {
        f"{W}before": str(round(paragraph.space_before_pt * 20)),
        f"{W}after": str(round(paragraph.space_after_pt * 20)),
    }
    if paragraph.line_spacing.type == "exact":
        spacing_attributes.update(
            {
                f"{W}line": str(round(paragraph.line_spacing.value_pt * 20)),
                f"{W}lineRule": "exact",
            }
        )
    else:
        spacing_attributes.update(
            {
                f"{W}line": str(round(paragraph.line_spacing.value * 240)),
                f"{W}lineRule": "auto",
            }
        )
    etree.SubElement(p_pr, f"{W}spacing", spacing_attributes)
    etree.SubElement(
        p_pr,
        f"{W}ind",
        {
            f"{W}firstLineChars": str(round(paragraph.first_line_chars * 100)),
            f"{W}left": str(round(paragraph.left_indent_mm * 1440 / 25.4)),
            f"{W}hanging": str(round(paragraph.hanging_indent_mm * 1440 / 25.4)),
        },
    )
    if paragraph.keep_with_next:
        etree.SubElement(p_pr, f"{W}keepNext")
    if paragraph.page_break_before:
        etree.SubElement(p_pr, f"{W}pageBreakBefore")

    run = expected.run
    r_pr = etree.SubElement(style, f"{W}rPr")
    etree.SubElement(
        r_pr,
        f"{W}rFonts",
        {
            f"{W}eastAsia": run.font_zh,
            f"{W}ascii": run.font_latin,
            f"{W}hAnsi": run.font_latin,
        },
    )
    etree.SubElement(r_pr, f"{W}sz", {f"{W}val": str(round(run.size_pt * 2))})
    etree.SubElement(r_pr, f"{W}b", {f"{W}val": "1" if run.bold else "0"})
    etree.SubElement(r_pr, f"{W}i", {f"{W}val": "1" if run.italic else "0"})
    _replace_part(path, "word/styles.xml", root)


def test_num_id_zero_unrelated_styleref_and_bullet_are_not_heading_numbering(
    tmp_path: Path,
) -> None:
    zero = build_docx(
        tmp_path / "zero.docx",
        [
            {
                "text": "1.1 Manual heading",
                "style": "AppleHeading2",
                "numbering_level": 1,
                "numbering_num_id": 0,
            }
        ],
    )
    styleref = build_docx(
        tmp_path / "styleref.docx",
        [
            {
                "text": "1.1 Manual heading",
                "style": "AppleHeading2",
                "field": " STYLEREF AppleBody ",
            }
        ],
    )
    bullet = build_docx(
        tmp_path / "bullet.docx",
        [
            {
                "text": "1.1 Manual heading",
                "style": "AppleHeading2",
                "numbering_level": 1,
            }
        ],
    )
    _set_num_format(bullet, "1", "bullet")

    assert inspect_document(zero, PROFILE).paragraphs[0].role == "heading_2"
    assert "NUMBERING.MANUAL_HEADING" in _rule_ids(zero)
    assert "NUMBERING.MANUAL_HEADING" in _rule_ids(styleref)
    assert "NUMBERING.MANUAL_HEADING" in _rule_ids(bullet)


@pytest.mark.parametrize(
    ("role", "style", "label"),
    [
        ("figure_caption", "AppleFigureCaption", "Figure"),
        ("table_caption", "AppleTableCaption", "Table"),
    ],
)
@pytest.mark.parametrize(
    ("has_styleref", "has_restarting_sequence", "is_automatic"),
    [
        (False, False, False),
        (True, False, False),
        (False, True, False),
        (True, True, True),
    ],
)
def test_caption_requires_both_chapter_styleref_and_restarting_sequence(
    tmp_path: Path,
    role: str,
    style: str,
    label: str,
    has_styleref: bool,
    has_restarting_sequence: bool,
    is_automatic: bool,
) -> None:
    fields = []
    if has_styleref:
        fields.append(" STYLEREF AppleHeading1 \\n ")
    if has_restarting_sequence:
        fields.append(f" SEQ {label} \\s 1 \\* ARABIC ")
    path = build_docx(
        tmp_path / f"{role}-{has_styleref}-{has_restarting_sequence}.docx",
        [
            {
                "text": "图 1-1 Hard coded" if label == "Figure" else "表 1-1 Hard coded",
                "style": style,
                "fields": fields,
            },
        ],
    )

    rule_ids = _rule_ids(path)
    assert ("NUMBERING.MANUAL_CAPTION" not in rule_ids) is is_automatic


def test_style_definition_audits_each_distinct_used_alias(tmp_path: Path) -> None:
    path = build_docx(
        tmp_path / "body-aliases.docx",
        [
            {"text": "First body paragraph", "style": "AppleBody"},
            {"text": "Second body paragraph", "style": "苹果大学正文"},
        ],
    )
    _set_style_to_profile(path, "AppleBody", "body")

    report = audit_document(path, PROFILE, clock=lambda: NOW)
    findings = [
        finding
        for finding in report["findings"]
        if finding["rule_id"] == "STYLE.DEFINITION"
    ]

    assert len(findings) == 1
    assert findings[0]["before"]["style_id"] == "苹果大学正文"


def test_table_caption_adjacency_uses_main_story_block_order(
    tmp_path: Path,
) -> None:
    path = build_docx(
        tmp_path / "multiple-tables.docx",
        [
            {
                "text": "表",
                "style": "AppleTableCaption",
                "fields": ["STYLEREF AppleHeading1", "SEQ Table \\s 1"],
            },
            {
                "text": "表",
                "style": "AppleTableCaption",
                "fields": ["STYLEREF AppleHeading1", "SEQ Table \\s 1"],
            },
        ],
        add_table=True,
    )
    root = _xml(path, "word/document.xml")
    body = root.find("w:body", namespaces=NS)
    assert body is not None
    captions = body.findall("w:p", namespaces=NS)
    table = body.find("w:tbl", namespaces=NS)
    assert len(captions) == 2
    assert table is not None
    body.remove(table)
    first_caption_index = body.index(captions[0])
    body.insert(first_caption_index + 1, table)
    second_caption_index = body.index(captions[1])
    body.insert(second_caption_index + 1, deepcopy(table))
    _replace_part(path, "word/document.xml", root)

    assert "TABLE.CAPTION.POSITION" not in _rule_ids(path)


def test_identical_missing_crossrefs_have_unique_deterministic_findings(
    tmp_path: Path,
) -> None:
    path = build_docx(
        tmp_path / "duplicate-crossrefs.docx",
        [
            {
                "text": "Repeated references",
                "fields": ["REF missing_target", "REF missing_target"],
            }
        ],
    )

    first = audit_document(path, PROFILE, clock=lambda: NOW)
    second = audit_document(path, PROFILE, clock=lambda: NOW)
    findings = [
        finding
        for finding in first["findings"]
        if finding["rule_id"] == "FIELD.CROSS_REFERENCE.TARGET"
    ]

    assert len(findings) == 2
    assert len({finding["finding_id"] for finding in findings}) == 2
    assert sorted(
        finding["before"]["field_occurrence"] for finding in findings
    ) == [1, 2]
    assert first["findings"] == second["findings"]


def test_three_line_table_requires_visible_top_header_bottom_and_bottom(
    tmp_path: Path,
) -> None:
    path = build_docx(
        tmp_path / "table.docx",
        [{"text": "表", "field": "SEQ Table \\s 1", "style": "AppleTableCaption"}],
        add_table=True,
    )
    root = _xml(path, "word/document.xml")
    table = root.find(".//w:tbl", namespaces=NS)
    assert table is not None
    tbl_pr = table.find("w:tblPr", namespaces=NS)
    assert tbl_pr is not None
    borders = etree.SubElement(tbl_pr, f"{W}tblBorders")
    for name, value in (("top", "nil"), ("bottom", "single")):
        etree.SubElement(borders, f"{W}{name}", {f"{W}val": value})
    first_row = table.find("w:tr", namespaces=NS)
    assert first_row is not None
    for cell in first_row.findall("w:tc", namespaces=NS):
        tc_pr = cell.find("w:tcPr", namespaces=NS)
        assert tc_pr is not None
        tc_borders = etree.SubElement(tc_pr, f"{W}tcBorders")
        etree.SubElement(tc_borders, f"{W}bottom", {f"{W}val": "single"})
    _replace_part(path, "word/document.xml", root)

    assert "TABLE.BORDERS.THREE_LINE" in _rule_ids(path)

    root = _xml(path, "word/document.xml")
    top = root.find(".//w:tbl/w:tblPr/w:tblBorders/w:top", namespaces=NS)
    assert top is not None
    top.set(f"{W}val", "single")
    _replace_part(path, "word/document.xml", root)
    assert "TABLE.BORDERS.THREE_LINE" not in _rule_ids(path)

    root = _xml(path, "word/document.xml")
    borders = root.find(".//w:tbl/w:tblPr/w:tblBorders", namespaces=NS)
    assert borders is not None
    etree.SubElement(borders, f"{W}insideV", {f"{W}val": "single"})
    _replace_part(path, "word/document.xml", root)
    assert "TABLE.BORDERS.THREE_LINE" in _rule_ids(path)


def test_reference_line_spacing_distinguishes_exact_from_auto(
    tmp_path: Path,
) -> None:
    exact = build_docx(
        tmp_path / "reference-exact.docx",
        [
            {"text": "参考文献"},
            {"text": "Reference entry"},
        ],
    )
    _set_paragraph_spacing(exact, 2, line=320, line_rule="exact")
    assert "REFERENCE.PARAGRAPH_FORMAT" not in _rule_ids(exact)

    auto = build_docx(
        tmp_path / "reference-auto.docx",
        [
            {"text": "参考文献"},
            {"text": "Reference entry"},
        ],
    )
    _set_paragraph_spacing(auto, 2, line=320, line_rule="auto")
    assert "REFERENCE.PARAGRAPH_FORMAT" in _rule_ids(auto)


def test_header_linking_uses_targets_and_dynamic_sources_require_styleref(
    tmp_path: Path,
) -> None:
    path = build_docx(
        tmp_path / "headers.docx",
        [
            {"text": "封面", "style": "AppleTitleZH", "section_break": True},
            {"text": "第一章 正文", "style": "AppleHeading1"},
        ],
        add_header_footer=True,
    )
    root = _xml(path, "word/document.xml")
    sections = root.xpath(".//w:sectPr", namespaces=NS)
    first_header = sections[0].find("w:headerReference", namespaces=NS)
    assert first_header is not None
    second_header = etree.SubElement(sections[1], f"{W}headerReference")
    second_header.set(f"{W}type", "default")
    second_header.set(f"{R}id", first_header.get(f"{R}id"))
    _replace_part(path, "word/document.xml", root)

    rule_ids = _rule_ids(path)
    assert "HEADER.FOOTER.LINKING" not in rule_ids
    assert "HEADER.FOOTER.HEADER_SOURCE" in rule_ids

    header = _xml(path, "word/header1.xml")
    paragraph = header.find(".//w:p", namespaces=NS)
    assert paragraph is not None
    for child in list(paragraph):
        paragraph.remove(child)
    field = etree.SubElement(paragraph, f"{W}fldSimple")
    field.set(f"{W}instr", " STYLEREF AppleHeading1 ")
    _replace_part(path, "word/header1.xml", header)
    assert "HEADER.FOOTER.HEADER_SOURCE" not in _rule_ids(path)

    unlinked = build_docx(
        tmp_path / "unlinked-headers.docx",
        [
            {"text": "封面", "style": "AppleTitleZH", "section_break": True},
            {"text": "第一章 正文", "style": "AppleHeading1"},
        ],
        add_header_footer=True,
    )
    document = Document(unlinked)
    document.sections[1].header.is_linked_to_previous = False
    document.sections[1].header.paragraphs[0].text = "Different header"
    document.save(unlinked)
    assert "HEADER.FOOTER.LINKING" in _rule_ids(unlinked)


def test_header_footer_nested_findings_use_part_paragraph_locations(
    tmp_path: Path,
) -> None:
    path = build_docx(
        tmp_path / "story-findings.docx",
        [{"text": "正文"}],
        add_header_footer=True,
    )
    for part_name, private_text, missing_target in (
        ("word/header1.xml", "author@example.com", "missing_header"),
        ("word/footer1.xml", "12345678", "missing_footer"),
    ):
        root = _xml(path, part_name)
        paragraph = root.find(".//w:p", namespaces=NS)
        assert paragraph is not None
        text = etree.SubElement(etree.SubElement(paragraph, f"{W}r"), f"{W}t")
        text.text = f" {private_text}"
        field = etree.SubElement(paragraph, f"{W}fldSimple")
        field.set(f"{W}instr", f" REF {missing_target} ")
        _replace_part(path, part_name, root)

    semantic = inspect_document(path, PROFILE)
    assert [node.location for node in semantic.headers] == [
        "part:word/header1.xml#paragraph:1"
    ]
    assert [node.location for node in semantic.footers] == [
        "part:word/footer1.xml#paragraph:1"
    ]

    first = audit_document(path, PROFILE, clock=lambda: NOW)
    second = audit_document(path, PROFILE, clock=lambda: NOW)
    story_findings = [
        finding
        for finding in first["findings"]
        if finding["rule_id"]
        in {
            "FIELD.CROSS_REFERENCE.TARGET",
            "PRIVACY.PATTERN.EMAIL",
            "PRIVACY.PATTERN.STUDENT_ID",
        }
    ]

    assert [
        (finding["location"], finding["rule_id"]) for finding in story_findings
    ] == [
        (
            "part:word/footer1.xml#paragraph:1",
            "FIELD.CROSS_REFERENCE.TARGET",
        ),
        ("part:word/footer1.xml#paragraph:1", "PRIVACY.PATTERN.STUDENT_ID"),
        (
            "part:word/header1.xml#paragraph:1",
            "FIELD.CROSS_REFERENCE.TARGET",
        ),
        ("part:word/header1.xml#paragraph:1", "PRIVACY.PATTERN.EMAIL"),
    ]
    assert first["findings"] == second["findings"]
    validate_audit(first)


def test_front_and_body_page_resets_follow_semantic_region_starts(
    tmp_path: Path,
) -> None:
    path = build_docx(
        tmp_path / "pagination.docx",
        [
            {"text": "中文题名", "style": "AppleTitleZH", "section_break": True},
            {"text": "中文摘要", "style": "AppleAbstractZH", "section_break": True},
            {"text": "第一章 正文", "style": "AppleHeading1"},
        ],
    )
    _set_section_page_number(path, 1, start=None, fmt=None)
    _set_section_page_number(path, 2, start=1, fmt="upperRoman")
    _set_section_page_number(path, 3, start=1, fmt="decimal")

    rule_ids = _rule_ids(path)
    assert "HEADER.FOOTER.PAGE_RESET" not in rule_ids
    assert "HEADER.FOOTER.PAGE_STYLE" not in rule_ids

    _set_section_page_number(path, 2, start=2, fmt="upperRoman")
    assert "HEADER.FOOTER.PAGE_RESET" in _rule_ids(path)


def test_markdown_escapes_pipes_backslashes_and_newlines() -> None:
    report = {
        "schema_version": "1.0",
        "tool_version": "0.1.0",
        "created_at": "2026-10-20T10:00:00Z",
        "input": "test.docx",
        "document_sha256": "0" * 64,
        "profile": "apple-university-thesis",
        "summary": {"errors": 0, "violations": 0, "warnings": 1},
        "findings": [
            {
                "finding_id": "0123456789abcdef",
                "rule_id": "TEST.ESCAPE",
                "severity": "warning",
                "location": "document",
                "message": "pipe | slash \\ line\nnext",
                "repairable": False,
                "before": "a|b\\c\nnext",
                "after": None,
                "repaired": False,
            }
        ],
    }
    report["findings"][0]["finding_id"] = _finding_digest(
        "1.0", report["findings"][0]
    )[:16]

    markdown = render_audit_markdown(report)

    assert "pipe \\| slash \\\\ line<br>next" in markdown
    assert "a\\|b\\\\c<br>next" in markdown
