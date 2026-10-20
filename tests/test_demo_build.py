from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from thesis_format.audit import audit_document
from thesis_format.demo.build import build_before, build_template, build_thesis
from thesis_format.demo.defects import DEFECTS
from thesis_format.ooxml import OoxmlPackage
from thesis_format.profile import load_profile


ROOT = Path(__file__).resolve().parents[1]
PROFILE_PATH = ROOT / "rules" / "apple-university.yaml"
PROFILE = load_profile(PROFILE_PATH)
REQUIRED_BOOKMARKS = {
    "demo_cover",
    "demo_toc",
    "demo_body",
    "demo_figure",
    "demo_table",
    "demo_references",
    "demo_appendix",
}


def _bookmark_names(path: Path) -> set[str]:
    root = OoxmlPackage.open(path).xml("word/document.xml")
    return {
        str(value)
        for value in root.xpath(
            ".//w:bookmarkStart/@w:name",
            namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
        )
    }


def _field_codes(path: Path) -> tuple[str, ...]:
    root = OoxmlPackage.open(path).xml("word/document.xml")
    return tuple(
        " ".join(str(value).split())
        for value in root.xpath(
            ".//w:instrText/text() | .//w:fldSimple/@w:instr",
            namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
        )
    )


def _rule_ids(path: Path) -> list[str]:
    report = audit_document(path, PROFILE)
    return [str(item["rule_id"]) for item in report["findings"]]


def test_template_is_clean_anonymous_profile_document(tmp_path: Path) -> None:
    output = build_template(tmp_path / "template.docx", PROFILE)
    package = OoxmlPackage.open(output)
    document = Document(output)

    assert output.is_file()
    assert REQUIRED_BOOKMARKS <= _bookmark_names(output)
    assert {
        'TOC \\o "1-3" \\h \\z \\u',
        'TOC \\h \\z \\c "图"',
        'TOC \\h \\z \\c "表"',
    } <= set(_field_codes(output))
    assert {style.name for style in document.styles} >= {
        PROFILE.styles[role].word_styles[0] for role in PROFILE.styles
    }
    assert all(
        round(section.page_width.mm) == 210
        and round(section.page_height.mm) == 297
        for section in document.sections
    )
    assert "word/comments.xml" not in package.parts
    assert not package.xml("word/document.xml").xpath(
        ".//w:ins | .//w:del",
        namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
    )
    core = package.xml("docProps/core.xml")
    text = " ".join(core.itertext())
    assert "艾苹" not in text
    assert "APPLE-0001" not in text


def test_full_thesis_contains_fields_images_tables_and_back_matter(
    tmp_path: Path,
) -> None:
    output = build_thesis(tmp_path / "thesis.docx", PROFILE)
    package = OoxmlPackage.open(output)
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "单株苹果树叶片数量的估算方法研究" in text
    assert "ABSTRACT" in text
    assert "关键词：" in text
    assert "参考文献" in text
    assert "附录 A" in text
    assert "致谢" in text
    assert len(document.tables) >= 2
    assert len(package.related_parts("word/document.xml", "image")) >= 2
    assert REQUIRED_BOOKMARKS <= _bookmark_names(output)
    assert len(_field_codes(output)) >= 3
    descriptions = package.xml("word/document.xml").xpath(
        ".//wp:docPr/@descr",
        namespaces={
            "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
        },
    )
    assert len(descriptions) >= 2
    assert all(str(description).strip() for description in descriptions)


def test_figure_object_is_kept_with_below_caption_only_when_needed(
    tmp_path: Path,
) -> None:
    output = build_thesis(tmp_path / "thesis.docx", PROFILE)
    document = Document(output)
    drawing_indexes = [
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph._p.xpath(".//w:drawing")
    ]

    assert drawing_indexes
    for index in drawing_indexes:
        figure = document.paragraphs[index]
        caption = document.paragraphs[index + 1]
        assert figure.paragraph_format.keep_with_next is True
        assert caption.style.name == PROFILE.styles.figure_caption.word_styles[0]
        next_is_source = (
            index + 2 < len(document.paragraphs)
            and document.paragraphs[index + 2].style.name
            == PROFILE.styles.source_note.word_styles[0]
        )
        if next_is_source:
            assert caption.paragraph_format.keep_with_next is True
        else:
            assert caption.paragraph_format.keep_with_next in {False, None}


def test_defect_registry_has_pure_one_concern_functions() -> None:
    assert tuple(DEFECTS) == (
        "wrong_paper_size",
        "wrong_body_margins",
        "wrong_section_start",
        "mixed_fonts",
        "manual_heading_numbers",
        "unreliable_heading_styles",
        "inconsistent_spacing",
        "missing_page_numbers",
        "broken_headers_footers",
        "malformed_toc",
        "caption_positions",
        "oversized_figure",
        "table_border_noise",
        "table_page_behavior",
        "reference_indent",
        "direct_formatting",
    )
    assert all(callable(function) for function in DEFECTS.values())


@pytest.mark.parametrize(
    ("name", "expected_rule"),
    (
        ("wrong_paper_size", "PAGE.PAPER.SIZE"),
        ("wrong_body_margins", "PAGE.MARGIN.BODY"),
        ("wrong_section_start", "PAGE.SECTION.START"),
        ("mixed_fonts", "STYLE.DEFINITION"),
        ("manual_heading_numbers", "NUMBERING.MANUAL_HEADING"),
        ("unreliable_heading_styles", "STYLE.ROLE.MAPPING"),
        ("inconsistent_spacing", "STYLE.DEFINITION"),
        ("missing_page_numbers", "HEADER.FOOTER.PAGE_NUMBER"),
        ("broken_headers_footers", "HEADER.FOOTER.PAGE_NUMBER"),
        ("malformed_toc", "FIELD.TOC.REQUIRED"),
        ("caption_positions", "FIGURE.CAPTION.POSITION"),
        ("oversized_figure", "FIGURE.DIMENSIONS.MAX_WIDTH"),
        ("table_border_noise", "TABLE.BORDERS.THREE_LINE"),
        ("table_page_behavior", "TABLE.ROW_SPLIT"),
        ("reference_indent", "REFERENCE.PARAGRAPH_FORMAT"),
        ("direct_formatting", "STYLE.DIRECT_FORMATTING"),
    ),
)
def test_each_defect_triggers_its_intended_rule(
    tmp_path: Path, name: str, expected_rule: str
) -> None:
    clean = build_thesis(tmp_path / f"{name}-clean.docx", PROFILE)
    broken = tmp_path / f"{name}.docx"

    DEFECTS[name](clean, broken)

    OoxmlPackage.open(broken)
    clean_rules = _rule_ids(clean)
    broken_rules = _rule_ids(broken)
    assert broken_rules.count(expected_rule) > clean_rules.count(expected_rule)


def test_before_document_is_readable_and_broadly_broken(tmp_path: Path) -> None:
    output = build_before(tmp_path / "before.docx", PROFILE)
    package = OoxmlPackage.open(output)
    report = audit_document(output, PROFILE)
    rule_ids = [str(item["rule_id"]) for item in report["findings"]]
    groups = {rule_id.split(".", 1)[0] for rule_id in rule_ids}

    assert package.xml("word/document.xml").tag == qn("w:document")
    assert Document(output).paragraphs
    assert len(rule_ids) >= 15
    assert len(groups) >= 7


def test_build_module_cli_generates_requested_outputs(tmp_path: Path) -> None:
    template = tmp_path / "cli-template.docx"
    before = tmp_path / "cli-before.docx"

    subprocess.run(
        [
            sys.executable,
            "-m",
            "thesis_format.demo.build",
            "--template",
            str(template),
            "--before",
            str(before),
        ],
        cwd=ROOT,
        check=True,
    )

    OoxmlPackage.open(template)
    OoxmlPackage.open(before)
