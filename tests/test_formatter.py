from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches
from lxml import etree
from PIL import Image

from thesis_format.audit import audit_document
from thesis_format.contracts import validate_transformation_result
from thesis_format.demo.build import build_template
from thesis_format.demo.defects import DEFECTS
from thesis_format.format import format_document
from thesis_format.ooxml import OoxmlPackage
from thesis_format.profile import load_profile
from thesis_format.snapshot import extract_snapshot

from helpers.docx_factory import build_docx


ROOT = Path(__file__).parents[1]
PROFILE = load_profile(ROOT / "rules" / "apple-university.yaml")
NOW = datetime(2026, 10, 20, 10, 0, tzinfo=UTC)
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _add_heading_numbering(
    path: Path,
    *,
    num_id: str,
    abstract_id: str,
    chapter_format: str = "chineseCounting",
) -> None:
    package = OoxmlPackage.open(path)
    numbering = package.xml("word/numbering.xml")
    abstract = etree.Element(
        f"{{{NS['w']}}}abstractNum",
        {f"{{{NS['w']}}}abstractNumId": abstract_id},
    )
    etree.SubElement(
        abstract,
        f"{{{NS['w']}}}multiLevelType",
        {f"{{{NS['w']}}}val": "multilevel"},
    )
    for level, (number_format, level_text, style_id) in enumerate(
        (
            (chapter_format, "第%1章", "AppleHeading1"),
            ("decimal", "%1.%2", "AppleHeading2"),
            ("decimal", "%1.%2.%3", "AppleHeading3"),
        )
    ):
        level_node = etree.SubElement(
            abstract,
            f"{{{NS['w']}}}lvl",
            {f"{{{NS['w']}}}ilvl": str(level)},
        )
        etree.SubElement(
            level_node,
            f"{{{NS['w']}}}start",
            {f"{{{NS['w']}}}val": "1"},
        )
        etree.SubElement(
            level_node,
            f"{{{NS['w']}}}numFmt",
            {f"{{{NS['w']}}}val": number_format},
        )
        etree.SubElement(
            level_node,
            f"{{{NS['w']}}}pStyle",
            {f"{{{NS['w']}}}val": style_id},
        )
        etree.SubElement(
            level_node,
            f"{{{NS['w']}}}suff",
            {f"{{{NS['w']}}}val": "space" if level == 0 else "nothing"},
        )
        etree.SubElement(
            level_node,
            f"{{{NS['w']}}}lvlText",
            {f"{{{NS['w']}}}val": level_text},
        )
        if level:
            etree.SubElement(
                level_node,
                f"{{{NS['w']}}}lvlRestart",
                {f"{{{NS['w']}}}val": "1"},
            )
    first_num = numbering.find("w:num", namespaces=NS)
    numbering.insert(
        numbering.index(first_num) if first_num is not None else len(numbering),
        abstract,
    )
    number = etree.Element(
        f"{{{NS['w']}}}num",
        {f"{{{NS['w']}}}numId": num_id},
    )
    etree.SubElement(
        number,
        f"{{{NS['w']}}}abstractNumId",
        {f"{{{NS['w']}}}val": abstract_id},
    )
    numbering.append(number)
    package.write(path, {"word/numbering.xml": numbering})


@pytest.mark.parametrize(
    ("defect_name", "repaired_rule"),
    (
        ("wrong_paper_size", "PAGE.PAPER.SIZE"),
        ("wrong_body_margins", "PAGE.MARGIN.BODY"),
        ("wrong_section_start", "PAGE.SECTION.START"),
        ("mixed_fonts", "STYLE.DEFINITION"),
        ("manual_heading_numbers", "NUMBERING.MANUAL_HEADING"),
        ("unreliable_heading_styles", "STYLE.ROLE.MAPPING"),
        ("inconsistent_spacing", "STYLE.DEFINITION"),
        ("missing_page_numbers", "HEADER.FOOTER.PAGE_NUMBER"),
        ("broken_headers_footers", "HEADER.FOOTER.HEADER_SOURCE"),
        ("malformed_toc", "FIELD.TOC.REQUIRED"),
        ("caption_positions", "FIGURE.CAPTION.POSITION"),
        ("oversized_figure", "FIGURE.DIMENSIONS.MAX_WIDTH"),
        ("table_border_noise", "TABLE.BORDERS.THREE_LINE"),
        ("table_page_behavior", "TABLE.REPEAT_HEADER"),
        ("reference_indent", "REFERENCE.PARAGRAPH_FORMAT"),
        ("direct_formatting", "STYLE.DIRECT_FORMATTING"),
    ),
)
def test_formatter_repairs_each_demo_defect_without_content_loss(
    tmp_path: Path,
    defect_name: str,
    repaired_rule: str,
) -> None:
    source = build_template(tmp_path / "source.docx", PROFILE)
    DEFECTS[defect_name](source, source)
    before = source.read_bytes()
    output = tmp_path / "formatted.docx"

    report = format_document(source, output, PROFILE, clock=lambda: NOW)

    validate_transformation_result(report)
    assert source.read_bytes() == before
    assert report["content_comparison"]["equivalent"] is True
    assert repaired_rule not in {
        finding["rule_id"] for finding in report["post_audit"]["findings"]
    }


def test_full_repair_is_stable_preserves_unsupported_parts_and_clears_mandatory_findings(
    tmp_path: Path,
) -> None:
    source = build_template(tmp_path / "source.docx", PROFILE)
    for defect in DEFECTS.values():
        defect(source, source)
    with ZipFile(source, "a") as archive:
        archive.writestr("customXml/unsupported.bin", b"\x00opaque\xff")
    source_bytes = source.read_bytes()
    output = tmp_path / "formatted.docx"

    report = format_document(source, output, PROFILE, clock=lambda: NOW)

    validate_transformation_result(report)
    assert source.read_bytes() == source_bytes
    with ZipFile(output) as archive:
        assert archive.read("customXml/unsupported.bin") == b"\x00opaque\xff"
    assert report["content_comparison"]["equivalent"] is True
    assert report["post_audit"]["summary"]["violations"] == 0
    assert report["post_audit"]["summary"]["errors"] == 0
    assert report["changes"] == sorted(
        report["changes"],
        key=lambda change: (
            _location_key(change["location"]),
            change["rule_id"],
            change["change_id"],
        ),
    )

    second = format_document(
        source,
        tmp_path / "formatted-again.docx",
        PROFILE,
        clock=lambda: NOW,
    )
    assert [
        {key: value for key, value in change.items() if key != "change_id"}
        for change in report["changes"]
    ] == [
        {key: value for key, value in change.items() if key != "change_id"}
        for change in second["changes"]
    ]


def test_ambiguous_paragraph_is_untouched_and_reported_as_skipped(
    tmp_path: Path,
) -> None:
    source = build_docx(
        tmp_path / "ambiguous.docx",
        [{"text": "可能是标题也可能是正文"}],
        unsupported_part=True,
    )
    before_root = OoxmlPackage.open(source).xml("word/document.xml")
    before_xml = before_root.xpath(
        "string(.//w:p[1])", namespaces=NS
    )

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    after_root = OoxmlPackage.open(tmp_path / "formatted.docx").xml(
        "word/document.xml"
    )
    assert after_root.xpath("string(.//w:p[1])", namespaces=NS) == before_xml
    assert any(
        change["status"] == "skipped"
        and change["rule_id"] == "SEMANTIC.AMBIGUOUS_ROLE"
        for change in report["changes"]
    )


def test_missing_cross_reference_target_is_not_invented(
    tmp_path: Path,
) -> None:
    source = build_docx(
        tmp_path / "crossref.docx",
        [{"text": "参见", "field": " REF missing_target "}],
    )

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    assert "missing_target" not in OoxmlPackage.open(
        tmp_path / "formatted.docx"
    ).xml("word/document.xml").xpath(
        ".//w:bookmarkStart/@w:name", namespaces=NS
    )
    assert any(
        finding["rule_id"] == "FIELD.CROSS_REFERENCE.TARGET"
        for finding in report["post_audit"]["findings"]
    )
    assert any(
        change["rule_id"] == "FIELD.CROSS_REFERENCE.TARGET"
        and change["status"] == "skipped"
        for change in report["changes"]
    )


def test_output_contains_profile_numbering_fields_tables_and_page_geometry(
    tmp_path: Path,
) -> None:
    source = build_template(tmp_path / "source.docx", PROFILE)
    for defect in DEFECTS.values():
        defect(source, source)

    output = tmp_path / "formatted.docx"
    format_document(source, output, PROFILE, clock=lambda: NOW)
    package = OoxmlPackage.open(output)
    document = package.xml("word/document.xml")
    numbering = package.xml("word/numbering.xml")
    styles = package.xml("word/styles.xml")

    assert document.xpath(
        "count(.//w:sectPr/w:pgSz[@w:w='11906'][@w:h='16838'])",
        namespaces=NS,
    ) == len(document.xpath(".//w:sectPr", namespaces=NS))
    assert numbering.xpath(
        ".//w:lvl[w:numFmt[@w:val='chineseCounting']]"
        "[w:lvlText[@w:val='第%1章']]",
        namespaces=NS,
    )
    assert document.xpath(
        ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
        "[.//w:instrText[contains(., 'STYLEREF')]]"
        "[.//w:instrText[contains(., 'SEQ')]]",
        namespaces=NS,
    )
    assert styles.xpath(
        ".//w:style[@w:styleId='AppleBody']"
        "/w:pPr/w:spacing[@w:lineRule='exact'][@w:line='400']",
        namespaces=NS,
    )
    table = document.xpath(".//w:tbl", namespaces=NS)[0]
    assert table.xpath("./w:tr[1]/w:trPr/w:tblHeader", namespaces=NS)
    assert not table.xpath(
        ".//w:tr[not(w:trPr/w:cantSplit)]", namespaces=NS
    )
    assert not table.xpath(
        "./w:tblPr/w:tblBorders/*"
        "[not(self::w:top or self::w:bottom)]",
        namespaces=NS,
    )


def test_table_caption_is_moved_above_with_bound_object_identity(
    tmp_path: Path,
) -> None:
    source = build_template(tmp_path / "source.docx", PROFILE)
    package = OoxmlPackage.open(source)
    root = package.xml("word/document.xml")
    body = root.find("w:body", namespaces=NS)
    table = body.find("w:tbl", namespaces=NS)
    caption = table.getprevious()
    body.remove(caption)
    body.insert(body.index(table) + 1, caption)
    package.write(source, {"word/document.xml": root})

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    assert report["content_comparison"]["equivalent"] is True
    change = next(
        change
        for change in report["changes"]
        if change["rule_id"] == "TABLE.CAPTION.POSITION"
    )
    assert change["operation"] == "move_caption"
    assert change["status"] == "applied"
    assert change["node_id"]
    assert len(change["object_occurrence_sha256"]) == 64
    assert change["before"] == {
        "node_id": change["node_id"],
        "object_occurrence_sha256": change["object_occurrence_sha256"],
        "position": "below",
    }
    assert change["after"] == {
        "node_id": change["node_id"],
        "object_occurrence_sha256": change["object_occurrence_sha256"],
        "position": "above",
    }


def _caption_test_document(
    path: Path,
    blocks: tuple[str, ...],
) -> Path:
    def append_field(paragraph, instruction: str, result: str) -> None:
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        code = OxmlElement("w:instrText")
        code.set(qn("xml:space"), "preserve")
        code.text = instruction
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = result
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend((begin, code, separate, text, end))

    def add_caption(label: str, title: str, style: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.style = style
        paragraph.add_run(f"{label} ")
        append_field(paragraph, ' STYLEREF "AppleHeading1" \\n ', "1")
        paragraph.add_run("-")
        append_field(paragraph, f" SEQ {label} \\s 1 \\* ARABIC ", "1")
        paragraph.add_run(f" {title}")

    image_path = path.with_suffix(".png")
    Image.new("RGB", (8, 8), "white").save(image_path)
    document = Document()
    for style_name in ("AppleFigureCaption", "AppleTableCaption", "AppleHeading1"):
        if style_name not in document.styles:
            document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    for block in blocks:
        if block == "figure":
            document.add_paragraph().add_run().add_picture(str(image_path))
        elif block == "two_images":
            paragraph = document.add_paragraph()
            paragraph.add_run().add_picture(str(image_path))
            paragraph.add_run().add_picture(str(image_path))
        elif block == "figure_caption":
            add_caption("图", "苹果树", "AppleFigureCaption")
        elif block == "table":
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "苹果"
            table.cell(0, 1).text = "梨"
        elif block == "table_caption":
            add_caption("表", "苹果数据", "AppleTableCaption")
        elif block == "separator":
            document.add_paragraph("分隔正文")
        else:
            raise AssertionError(f"unknown block: {block}")
    document.save(path)
    image_path.unlink()
    return path


@pytest.mark.parametrize(
    ("blocks", "rule_id"),
    [
        (("figure", "figure", "figure_caption"), "FIGURE.CAPTION.POSITION"),
        (("table", "table", "table_caption"), "TABLE.CAPTION.POSITION"),
    ],
)
def test_ambiguous_single_caption_for_two_objects_is_not_moved(
    tmp_path: Path,
    blocks: tuple[str, ...],
    rule_id: str,
) -> None:
    source = _caption_test_document(tmp_path / "ambiguous.docx", blocks)
    before_snapshot = extract_snapshot(source, PROFILE)

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    after_snapshot = extract_snapshot(tmp_path / "formatted.docx", PROFILE)
    before_relevant = [
        (node["node_type"], node["node_id"])
        for node in before_snapshot["body"]
        if node["node_type"] in {"caption", "image", "table"}
    ]
    after_relevant = [
        (node["node_type"], node["node_id"])
        for node in after_snapshot["body"]
        if node["node_type"] in {"caption", "image", "table"}
    ]
    assert after_relevant == before_relevant
    changes = [
        change for change in report["changes"] if change["rule_id"] == rule_id
    ]
    assert changes
    assert all(change["status"] == "skipped" for change in changes)
    assert all(
        change["operation"] == "preserve-ambiguous-caption-association"
        for change in changes
    )
    assert any(
        finding["rule_id"] == rule_id
        for finding in report["post_audit"]["findings"]
    )


@pytest.mark.parametrize(
    ("blocks", "rule_id", "expected_position"),
    [
        (
            ("figure_caption", "figure"),
            "FIGURE.CAPTION.POSITION",
            "below",
        ),
        (
            ("table", "table_caption"),
            "TABLE.CAPTION.POSITION",
            "above",
        ),
    ],
)
def test_unique_swapped_caption_is_repaired_with_exact_occurrence_binding(
    tmp_path: Path,
    blocks: tuple[str, ...],
    rule_id: str,
    expected_position: str,
) -> None:
    source = _caption_test_document(tmp_path / "swapped.docx", blocks)

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    change = next(
        change for change in report["changes"] if change["rule_id"] == rule_id
    )
    assert change["status"] == "applied"
    assert change["operation"] == "move_caption"
    assert change["before"]["node_id"] == change["node_id"]
    assert (
        change["before"]["object_occurrence_sha256"]
        == change["object_occurrence_sha256"]
    )
    assert change["after"] == {
        **change["before"],
        "position": expected_position,
    }
    assert report["content_comparison"]["equivalent"] is True
    assert rule_id not in {
        finding["rule_id"] for finding in report["post_audit"]["findings"]
    }


def test_multi_image_paragraph_caption_is_not_moved_when_occurrence_is_ambiguous(
    tmp_path: Path,
) -> None:
    source = _caption_test_document(
        tmp_path / "multi-image.docx",
        ("figure_caption", "two_images"),
    )
    before = extract_snapshot(source, PROFILE)

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    after = extract_snapshot(tmp_path / "formatted.docx", PROFILE)
    assert [
        (node["node_type"], node["node_id"]) for node in after["body"]
    ] == [
        (node["node_type"], node["node_id"]) for node in before["body"]
    ]
    change = next(
        change
        for change in report["changes"]
        if change["rule_id"] == "FIGURE.CAPTION.POSITION"
    )
    assert change["status"] == "skipped"
    assert change["operation"] == "preserve-ambiguous-caption-association"


def test_second_table_caption_uses_second_table_occurrence_hash(
    tmp_path: Path,
) -> None:
    source = _caption_test_document(
        tmp_path / "second-table.docx",
        (
            "table_caption",
            "table",
            "separator",
            "table",
            "table_caption",
        ),
    )
    table_hashes = list(
        dict.fromkeys(
            str(node["relationship_hash"])
            for node in extract_snapshot(source, PROFILE)["body"]
            if node["node_type"] == "table"
        )
    )
    assert len(table_hashes) == 2

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    change = next(
        change
        for change in report["changes"]
        if change["rule_id"] == "TABLE.CAPTION.POSITION"
        and change["location"] == "table:2"
    )
    assert change["status"] == "applied"
    assert change["object_occurrence_sha256"] == table_hashes[1]


def test_numbering_repair_reuses_compatible_ids_without_changing_lists(
    tmp_path: Path,
) -> None:
    source = build_template(tmp_path / "numbering-collision.docx", PROFILE)
    DEFECTS["manual_heading_numbers"](source, source)
    package = OoxmlPackage.open(source)
    numbering = package.xml("word/numbering.xml")
    for identifier, level_text in (("910", "•"), ("911", "%1)")):
        abstract = etree.Element(
            f"{{{NS['w']}}}abstractNum",
            {f"{{{NS['w']}}}abstractNumId": identifier},
        )
        level = etree.SubElement(
            abstract,
            f"{{{NS['w']}}}lvl",
            {f"{{{NS['w']}}}ilvl": "0"},
        )
        etree.SubElement(
            level, f"{{{NS['w']}}}start", {f"{{{NS['w']}}}val": "1"}
        )
        etree.SubElement(
            level,
            f"{{{NS['w']}}}numFmt",
            {f"{{{NS['w']}}}val": "bullet" if identifier == "910" else "decimal"},
        )
        etree.SubElement(
            level,
            f"{{{NS['w']}}}lvlText",
            {f"{{{NS['w']}}}val": level_text},
        )
        num = etree.Element(
            f"{{{NS['w']}}}num",
            {f"{{{NS['w']}}}numId": identifier},
        )
        etree.SubElement(
            num,
            f"{{{NS['w']}}}abstractNumId",
            {f"{{{NS['w']}}}val": identifier},
        )
        first_num = numbering.find("w:num", namespaces=NS)
        numbering.insert(
            numbering.index(first_num) if first_num is not None else len(numbering),
            abstract,
        )
        numbering.append(num)
    protected = {
        identifier: (
            etree.tostring(
                numbering.xpath(
                    f"./w:abstractNum[@w:abstractNumId='{identifier}']",
                    namespaces=NS,
                )[0],
                method="c14n",
            ),
            etree.tostring(
                numbering.xpath(
                    f"./w:num[@w:numId='{identifier}']", namespaces=NS
                )[0],
                method="c14n",
            ),
        )
        for identifier in ("910", "911")
    }
    existing_heading_num_ids = package.xml("word/document.xml").xpath(
        ".//w:p[w:pPr/w:pStyle["
        "@w:val='AppleHeading1' or @w:val='AppleHeading2' "
        "or @w:val='AppleHeading3']]"
        "/w:pPr/w:numPr/w:numId/@w:val",
        namespaces=NS,
    )
    assert len(set(existing_heading_num_ids)) == 1
    existing_heading_num_id = existing_heading_num_ids[0]
    package.write(source, {"word/numbering.xml": numbering})

    output = tmp_path / "formatted.docx"
    format_document(source, output, PROFILE, clock=lambda: NOW)

    result = OoxmlPackage.open(output)
    result_numbering = result.xml("word/numbering.xml")
    for identifier, expected in protected.items():
        assert (
            etree.tostring(
                result_numbering.xpath(
                    f"./w:abstractNum[@w:abstractNumId='{identifier}']",
                    namespaces=NS,
                )[0],
                method="c14n",
            ),
            etree.tostring(
                result_numbering.xpath(
                    f"./w:num[@w:numId='{identifier}']", namespaces=NS
                )[0],
                method="c14n",
            ),
        ) == expected
    heading_num_id = result.xml("word/document.xml").xpath(
        "string(.//w:p[w:pPr/w:pStyle[@w:val='AppleHeading1']]"
        "[1]/w:pPr/w:numPr/w:numId/@w:val)",
        namespaces=NS,
    )
    assert heading_num_id == existing_heading_num_id
    abstract_id = result_numbering.xpath(
        "string(./w:num[@w:numId=$num_id]/w:abstractNumId/@w:val)",
        namespaces=NS,
        num_id=heading_num_id,
    )
    assert not result_numbering.xpath(
        "./w:abstractNum[@w:abstractNumId=$abstract_id]"
        "[w:nsid[@w:val='54464844']][w:tmpl[@w:val='48454144']]",
        namespaces=NS,
        abstract_id=abstract_id,
    )


def test_manual_chapter_reuses_existing_automatic_heading_sequence(
    tmp_path: Path,
) -> None:
    source = build_docx(
        tmp_path / "shared-heading-sequence.docx",
        [
            {
                "text": "绪论",
                "style": "AppleHeading1",
                "numbering_level": 0,
                "numbering_num_id": 901,
                "outline_level": 0,
            },
            {"text": "第二章 方法", "style": "AppleHeading1"},
        ],
    )
    _add_heading_numbering(source, num_id="901", abstract_id="901")

    output = tmp_path / "formatted.docx"
    format_document(source, output, PROFILE, clock=lambda: NOW)

    document = OoxmlPackage.open(output).xml("word/document.xml")
    heading_num_ids = document.xpath(
        ".//w:p[w:pPr/w:pStyle[@w:val='AppleHeading1']]"
        "/w:pPr/w:numPr/w:numId/@w:val",
        namespaces=NS,
    )
    assert heading_num_ids == ["901", "901"]
    assert not OoxmlPackage.open(output).xml("word/numbering.xml").xpath(
        "./w:abstractNum[w:nsid[@w:val='54464844']]",
        namespaces=NS,
    )


def test_manual_caption_reuses_existing_accepted_figure_sequence_identifier(
    tmp_path: Path,
) -> None:
    source = build_docx(
        tmp_path / "shared-caption-sequence.docx",
        [
            {
                "text": "图 苹果树",
                "style": "AppleFigureCaption",
                "field": " SEQ Figure \\s 1 \\* ARABIC ",
            },
            {"text": "图 1-2 梨树", "style": "AppleFigureCaption"},
        ],
    )

    output = tmp_path / "formatted.docx"
    format_document(source, output, PROFILE, clock=lambda: NOW)

    document = OoxmlPackage.open(output).xml("word/document.xml")
    sequence_fields = [
        " ".join(value.split())
        for value in document.xpath(
            ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
            "//w:instrText/text()",
            namespaces=NS,
        )
        if value.strip().upper().startswith("SEQ ")
    ]
    assert len(sequence_fields) == 2
    assert all(field.startswith("SEQ Figure ") for field in sequence_fields)
    assert not any(field.startswith("SEQ 图 ") for field in sequence_fields)


@pytest.mark.parametrize("representation", ("split", "simple"))
def test_manual_caption_reuses_accepted_identifier_from_valid_field_forms(
    tmp_path: Path,
    representation: str,
) -> None:
    source = build_docx(
        tmp_path / f"{representation}-caption-sequence.docx",
        [
            {
                "text": "图 苹果树",
                "style": "AppleFigureCaption",
                "field": " SEQ Figure \\s 1 \\* ARABIC ",
            },
            {"text": "图 1-2 梨树", "style": "AppleFigureCaption"},
        ],
    )
    package = OoxmlPackage.open(source)
    document = package.xml("word/document.xml")
    first_caption = document.xpath(
        ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']][1]",
        namespaces=NS,
    )[0]
    instruction = first_caption.find(".//w:instrText", namespaces=NS)
    assert instruction is not None
    if representation == "split":
        original = instruction.text or ""
        instruction.text = " SEQ "
        continuation = etree.Element(f"{{{NS['w']}}}instrText")
        continuation.text = original[len(" SEQ ") :]
        instruction.addnext(continuation)
    else:
        field_run = instruction.getparent()
        assert field_run is not None
        first_caption.remove(field_run)
        simple = etree.Element(
            f"{{{NS['w']}}}fldSimple",
            {f"{{{NS['w']}}}instr": " SEQ Figure \\s 1 \\* ARABIC "},
        )
        run = etree.SubElement(simple, f"{{{NS['w']}}}r")
        etree.SubElement(run, f"{{{NS['w']}}}t").text = "1"
        first_caption.append(simple)
    package.write(source, {"word/document.xml": document})

    output = tmp_path / "formatted.docx"
    format_document(source, output, PROFILE, clock=lambda: NOW)

    result = OoxmlPackage.open(output).xml("word/document.xml")
    field_text = " ".join(
        result.xpath(
            "string(.//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
            "[2]//w:instrText[contains(., 'SEQ')])",
            namespaces=NS,
        ).split()
    )
    assert field_text.startswith("SEQ Figure ")


def test_manual_caption_skips_incompatible_existing_sequence_identifier(
    tmp_path: Path,
) -> None:
    source = build_docx(
        tmp_path / "incompatible-caption-sequence.docx",
        [
            {
                "text": "图 苹果树",
                "style": "AppleFigureCaption",
                "field": " SEQ Figure \\s 1 \\* ARABIC ",
            },
            {
                "text": "图 梨树",
                "style": "AppleFigureCaption",
                "field": " SEQ Fig \\s 1 \\* ARABIC ",
            },
            {"text": "图 1-3 桃树", "style": "AppleFigureCaption"},
        ],
    )

    output = tmp_path / "formatted.docx"
    with pytest.warns(
        RuntimeWarning,
        match="incompatible figure caption SEQ identifiers",
    ):
        report = format_document(source, output, PROFILE, clock=lambda: NOW)

    result = OoxmlPackage.open(output).xml("word/document.xml")
    captions = result.xpath(
        ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]",
        namespaces=NS,
    )
    assert "图 1-3" in "".join(captions[2].xpath(".//w:t/text()", namespaces=NS))
    assert not captions[2].xpath(".//w:instrText", namespaces=NS)
    assert any(
        change["rule_id"] == "NUMBERING.MANUAL_CAPTION"
        and change["status"] == "skipped"
        for change in report["changes"]
    )


def test_manual_heading_skips_incompatible_existing_numbering_sequences(
    tmp_path: Path,
) -> None:
    source = build_docx(
        tmp_path / "mixed-heading-sequences.docx",
        [
            {
                "text": "绪论",
                "style": "AppleHeading1",
                "numbering_level": 0,
                "numbering_num_id": 901,
                "outline_level": 0,
            },
            {
                "text": "方法",
                "style": "AppleHeading1",
                "numbering_level": 0,
                "numbering_num_id": 902,
                "outline_level": 0,
            },
            {"text": "第三章 结果", "style": "AppleHeading1"},
        ],
    )
    _add_heading_numbering(source, num_id="901", abstract_id="901")
    _add_heading_numbering(
        source,
        num_id="902",
        abstract_id="902",
        chapter_format="decimal",
    )

    output = tmp_path / "formatted.docx"
    with pytest.warns(RuntimeWarning, match="incompatible heading numbering"):
        report = format_document(source, output, PROFILE, clock=lambda: NOW)

    document = OoxmlPackage.open(output).xml("word/document.xml")
    manual = document.xpath(
        ".//w:p[w:pPr/w:pStyle[@w:val='AppleHeading1']][3]",
        namespaces=NS,
    )[0]
    assert not manual.xpath("./w:pPr/w:numPr", namespaces=NS)
    assert "第三章" in "".join(manual.xpath(".//w:t/text()", namespaces=NS))
    assert any(
        change["rule_id"] == "NUMBERING.MANUAL_HEADING"
        and change["status"] == "skipped"
        for change in report["changes"]
    )


def test_caption_repair_preserves_structured_content_and_snapshot(
    tmp_path: Path,
) -> None:
    source = build_template(tmp_path / "structured-caption.docx", PROFILE)
    image_path = tmp_path / "caption.png"
    Image.new("RGB", (8, 8), "blue").save(image_path)
    document = Document(source)
    caption = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "AppleFigureCaption"
    )
    caption.clear()
    token = caption.add_run("图 1-1")
    document.add_comment(token, "Keep this comment", author="Reviewer")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(
        qn("r:id"),
        caption.part.relate_to(
            "https://example.com/apple",
            RELATIONSHIP_TYPE.HYPERLINK,
            is_external=True,
        ),
    )
    hyperlink_run = OxmlElement("w:r")
    hyperlink_text = OxmlElement("w:t")
    hyperlink_text.set(qn("xml:space"), "preserve")
    hyperlink_text.text = " 苹果树"
    hyperlink_run.append(hyperlink_text)
    hyperlink.append(hyperlink_run)
    caption._p.append(hyperlink)
    _append_test_field(caption, " DATE ")
    note_run = caption.add_run()
    note_reference = OxmlElement("w:footnoteReference")
    note_reference.set(qn("w:id"), "9")
    note_run._r.append(note_reference)
    caption.add_run().add_picture(str(image_path), width=Inches(0.1))
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), "77")
    bookmark_start.set(qn("w:name"), "structured_caption")
    caption._p.insert(1, bookmark_start)
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), "77")
    caption._p.append(bookmark_end)
    document.save(source)
    before_snapshot = extract_snapshot(source, PROFILE)
    before = OoxmlPackage.open(source).xml("word/document.xml")
    protected_before = {
        name: [
            etree.tostring(element, method="c14n")
            for element in before.xpath(xpath, namespaces=NS)
        ]
        for name, xpath in {
            "hyperlinks": ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
            "/w:hyperlink",
            "note_refs": ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
            "//w:footnoteReference/parent::w:r",
            "comments": ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
            "/*[self::w:commentRangeStart or self::w:commentRangeEnd]"
            " | .//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
            "//w:commentReference/parent::w:r",
            "drawings": ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
            "//w:drawing",
            "bookmarks": ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
            "/*[self::w:bookmarkStart or self::w:bookmarkEnd]",
        }.items()
    }

    output = tmp_path / "formatted.docx"
    report = format_document(source, output, PROFILE, clock=lambda: NOW)

    after_snapshot = extract_snapshot(output, PROFILE)
    assert report["content_comparison"]["equivalent"] is True, report[
        "content_comparison"
    ]
    assert before_snapshot["body"] == after_snapshot["body"]
    after = OoxmlPackage.open(output).xml("word/document.xml")
    for name, xpath in {
        "hyperlinks": ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
        "/w:hyperlink",
        "note_refs": ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
        "//w:footnoteReference/parent::w:r",
        "comments": ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
        "/*[self::w:commentRangeStart or self::w:commentRangeEnd]"
        " | .//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
        "//w:commentReference/parent::w:r",
        "drawings": ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
        "//w:drawing",
        "bookmarks": ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
        "/*[self::w:bookmarkStart or self::w:bookmarkEnd]",
    }.items():
        assert [
            etree.tostring(element, method="c14n")
            for element in after.xpath(xpath, namespaces=NS)
        ] == protected_before[name]
    assert after.xpath(
        ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
        "//w:instrText[normalize-space(.)='DATE']",
        namespaces=NS,
    )
    assert after.xpath(
        ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]"
        "[.//w:instrText[contains(., 'STYLEREF')]]"
        "[.//w:instrText[contains(., 'SEQ')]]",
        namespaces=NS,
    )


def test_caption_repair_skips_manual_number_inside_hyperlink(
    tmp_path: Path,
) -> None:
    source = build_template(tmp_path / "unsafe-caption.docx", PROFILE)
    document = Document(source)
    caption = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "AppleFigureCaption"
    )
    caption.clear()
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(
        qn("r:id"),
        caption.part.relate_to(
            "https://example.com/number",
            RELATIONSHIP_TYPE.HYPERLINK,
            is_external=True,
        ),
    )
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "图 1-1"
    run.append(text)
    hyperlink.append(run)
    caption._p.append(hyperlink)
    caption.add_run(" 苹果树")
    document.save(source)
    before = OoxmlPackage.open(source).xml("word/document.xml")
    before_caption = etree.tostring(
        before.xpath(
            ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]",
            namespaces=NS,
        )[0],
        method="c14n",
    )

    output = tmp_path / "formatted.docx"
    report = format_document(source, output, PROFILE, clock=lambda: NOW)

    after = OoxmlPackage.open(output).xml("word/document.xml")
    assert etree.tostring(
        after.xpath(
            ".//w:p[w:pPr/w:pStyle[@w:val='AppleFigureCaption']]",
            namespaces=NS,
        )[0],
        method="c14n",
    ) == before_caption
    assert any(
        change["rule_id"] == "NUMBERING.MANUAL_CAPTION"
        and change["status"] == "skipped"
        for change in report["changes"]
    )
    assert any(
        finding["rule_id"] == "NUMBERING.MANUAL_CAPTION"
        for finding in report["post_audit"]["findings"]
    )


def test_page_number_repair_preserves_footer_text_table_image_and_field(
    tmp_path: Path,
) -> None:
    source = build_docx(
        tmp_path / "footer-content.docx",
        [{"text": "第一章 绪论", "style": "AppleHeading1"}],
    )
    image_path = tmp_path / "footer.png"
    Image.new("RGB", (8, 8), "green").save(image_path)
    document = Document(source)
    footer = document.sections[0].footer
    footer.paragraphs[0].text = "Confidential footer"
    _append_test_field(footer.add_paragraph(), " DATE ")
    footer.add_paragraph().add_run().add_picture(
        str(image_path), width=Inches(0.1)
    )
    footer_table = footer.add_table(rows=1, cols=1, width=Inches(1))
    footer_table.cell(0, 0).text = "Footer table"
    document.save(source)

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    package = OoxmlPackage.open(tmp_path / "formatted.docx")
    footer_part = next(iter(package.iter_part_names("word/footer")))
    root = package.xml(footer_part)
    assert root.xpath(
        ".//w:p[.//w:t='Confidential footer']", namespaces=NS
    )
    assert root.xpath(".//w:tbl[.//w:t='Footer table']", namespaces=NS)
    assert root.xpath(".//w:drawing", namespaces=NS)
    assert root.xpath(
        ".//w:instrText[contains(., 'DATE')]", namespaces=NS
    )
    page_paragraphs = root.xpath(
        "./w:p[.//w:instrText[contains(., 'PAGE')]]", namespaces=NS
    )
    assert len(page_paragraphs) == 1
    assert page_paragraphs[0].xpath(
        "./w:pPr/w:jc[@w:val='center']", namespaces=NS
    )
    assert report["content_comparison"]["equivalent"] is True


def test_page_number_exemption_does_not_hide_footer_content_loss(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = build_docx(
        tmp_path / "footer-loss.docx",
        [{"text": "第一章 绪论", "style": "AppleHeading1"}],
    )
    document = Document(source)
    document.sections[0].footer.paragraphs[0].text = "Must survive"
    document.save(source)

    from thesis_format.format import headers

    original = headers._write_footer

    def destructive_footer(root, alignment):
        original(root, alignment)
        text = root.xpath(
            "./w:p//w:t[.='Must survive']", namespaces=NS
        )[0]
        text.getparent().getparent().remove(text.getparent())

    monkeypatch.setattr(headers, "_write_footer", destructive_footer)

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    assert report["content_comparison"]["equivalent"] is False
    assert any(
        difference["kind"] == "removed"
        and difference["before"]["normalized_text"] == "Must survive"
        for difference in report["content_comparison"]["differences"]
    )


def test_table_repairs_ignore_nested_tables_and_target_second_outer_table(
    tmp_path: Path,
) -> None:
    source = build_docx(
        tmp_path / "nested-tables.docx",
        [{"text": "第一章 绪论", "style": "AppleHeading1"}],
    )
    document = Document(source)
    if "AppleTableCaption" not in document.styles:
        document.styles.add_style(
            "AppleTableCaption", WD_STYLE_TYPE.PARAGRAPH
        )
    document.add_paragraph("表 1-1 First outer", style="AppleTableCaption")
    first_outer = document.add_table(rows=1, cols=1)
    first_outer.cell(0, 0).text = "Outer one"
    nested = first_outer.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested sentinel"
    document.add_paragraph("Between tables")
    second_outer = document.add_table(rows=1, cols=1)
    second_outer.cell(0, 0).text = "Outer two"
    document.add_paragraph("表 1-2 Second outer", style="AppleTableCaption")
    document.save(source)

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    package = OoxmlPackage.open(tmp_path / "formatted.docx")
    body = package.xml("word/document.xml").find("w:body", namespaces=NS)
    assert body is not None
    outer_tables = body.findall("w:tbl", namespaces=NS)
    assert len(outer_tables) == 2
    assert outer_tables[1].xpath(
        "./w:tr[1]/w:trPr/w:tblHeader", namespaces=NS
    )
    assert outer_tables[1].xpath(
        "./w:tblPr/w:tblBorders/w:top", namespaces=NS
    )
    output_nested = outer_tables[0].xpath(
        "./w:tr/w:tc/w:tbl", namespaces=NS
    )[0]
    assert not output_nested.xpath(
        "./w:tr[1]/w:trPr/w:tblHeader", namespaces=NS
    )
    second_caption = outer_tables[1].getprevious()
    assert second_caption is not None
    assert "Second outer" in "".join(
        second_caption.xpath(".//w:t/text()", namespaces=NS)
    )
    change = next(
        change
        for change in report["changes"]
        if change["rule_id"] == "TABLE.CAPTION.POSITION"
        and change["location"] == "table:2"
    )
    assert change["status"] == "applied"
    assert not any(
        str(change["location"]) == "table:3"
        for change in report["changes"]
    )
    top_level_hashes = list(
        dict.fromkeys(
            str(node["relationship_hash"])
            for node in extract_snapshot(source, PROFILE)["body"]
            if node["node_type"] == "table"
            and str(node["source_location"]["path"]).startswith(
                "/w:document/w:body/w:tbl["
            )
        )
    )
    assert change["object_occurrence_sha256"] == top_level_hashes[1]


def test_formatter_creates_missing_body_header_footer_package_parts(
    tmp_path: Path,
) -> None:
    source = build_docx(
        tmp_path / "blank-stories.docx",
        [{"text": "第一章 绪论", "style": "AppleHeading1"}],
    )
    before = OoxmlPackage.open(source)
    assert not tuple(before.iter_part_names("word/header"))
    assert not tuple(before.iter_part_names("word/footer"))

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    package = OoxmlPackage.open(tmp_path / "formatted.docx")
    document = package.xml("word/document.xml")
    relationships = package.xml("word/_rels/document.xml.rels")
    content_types = package.xml("[Content_Types].xml")
    header_references = document.xpath(
        ".//w:sectPr/w:headerReference[@w:type='default']/@r:id",
        namespaces=NS,
    )
    footer_references = document.xpath(
        ".//w:sectPr/w:footerReference[@w:type='default']/@r:id",
        namespaces=NS,
    )
    assert len(header_references) == len(footer_references) == 1
    targets = {
        relationship.get("Id"): relationship.get("Target")
        for relationship in relationships.findall(f"{{{REL_NS}}}Relationship")
    }
    header_part = f"word/{targets[header_references[0]]}"
    footer_part = f"word/{targets[footer_references[0]]}"
    assert header_part in package.parts
    assert footer_part in package.parts
    assert package.xml(header_part).xpath(
        ".//w:instrText[contains(., 'STYLEREF')]", namespaces=NS
    )
    assert package.xml(footer_part).xpath(
        ".//w:instrText[contains(., 'PAGE')]", namespaces=NS
    )
    assert package.xml(footer_part).xpath(
        ".//w:pPr/w:jc[@w:val='center']", namespaces=NS
    )
    overrides = {
        item.get("PartName"): item.get("ContentType")
        for item in content_types.findall(
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override"
        )
    }
    assert overrides[f"/{header_part}"].endswith("header+xml")
    assert overrides[f"/{footer_part}"].endswith("footer+xml")
    corresponding = {
        "HEADER.FOOTER.HEADER_REQUIRED",
        "HEADER.FOOTER.HEADER_SOURCE",
        "HEADER.FOOTER.FOOTER_REQUIRED",
        "HEADER.FOOTER.PAGE_NUMBER",
        "HEADER.FOOTER.PAGE_ALIGNMENT",
    }
    assert not (
        corresponding
        & {
            finding["rule_id"]
            for finding in report["post_audit"]["findings"]
        }
    )
    assert all(
        change["status"] == "applied"
        for change in report["changes"]
        if change["rule_id"] in corresponding
    )


def test_formatter_does_not_create_stories_for_unnumbered_title_section(
    tmp_path: Path,
) -> None:
    source = build_docx(
        tmp_path / "title-only.docx",
        [{"text": "苹果大学学位论文", "style": "AppleTitleZH"}],
    )

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        PROFILE,
        clock=lambda: NOW,
    )

    package = OoxmlPackage.open(tmp_path / "formatted.docx")
    document = package.xml("word/document.xml")
    assert not document.xpath(
        ".//w:sectPr/w:headerReference | .//w:sectPr/w:footerReference",
        namespaces=NS,
    )
    assert not tuple(package.iter_part_names("word/header"))
    assert not tuple(package.iter_part_names("word/footer"))
    assert not {
        "HEADER.FOOTER.HEADER_REQUIRED",
        "HEADER.FOOTER.FOOTER_REQUIRED",
        "HEADER.FOOTER.PAGE_NUMBER",
    } & {
        finding["rule_id"] for finding in report["post_audit"]["findings"]
    }


def test_formatter_creates_fixed_text_header_from_profile(
    tmp_path: Path,
) -> None:
    body = replace(
        PROFILE.headers_footers.body,
        header_source="fixed_text",
        header_text="苹果大学",
    )
    profile = replace(
        PROFILE,
        headers_footers=replace(PROFILE.headers_footers, body=body),
    )
    source = build_docx(
        tmp_path / "fixed-header.docx",
        [{"text": "第一章 绪论", "style": "AppleHeading1"}],
    )

    report = format_document(
        source,
        tmp_path / "formatted.docx",
        profile,
        clock=lambda: NOW,
    )

    package = OoxmlPackage.open(tmp_path / "formatted.docx")
    header_part = next(
        name
        for name in package.parts
        if name.startswith("word/header") and name.endswith(".xml")
    )
    assert package.xml(header_part).xpath(
        "string(.//w:p)", namespaces=NS
    ) == "苹果大学"
    assert "HEADER.FOOTER.HEADER_SOURCE" not in {
        finding["rule_id"] for finding in report["post_audit"]["findings"]
    }


def _location_key(location: str) -> tuple[object, ...]:
    order = {
        "document": 0,
        "section": 1,
        "paragraph": 2,
        "table": 3,
        "cell": 4,
        "header": 5,
        "footer": 6,
        "part": 7,
    }
    if location == "document":
        return (0,)
    kind, value = location.split(":", 1)
    if kind == "cell":
        return (order[kind], *(int(item) for item in value.split("/")))
    if kind == "part":
        return (order[kind], value)
    return (order[kind], int(value))


def _append_test_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    for kind, text in (
        ("begin", None),
        (None, instruction),
        ("separate", None),
        ("end", None),
    ):
        if kind is not None:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            node = OxmlElement("w:instrText")
            node.text = text
        run._r.append(node)
