from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _append_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


def _append_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _set_numbering(paragraph, level: int = 0, num_id: int = 1) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num))
    p_pr.append(num_pr)


def _set_outline_level(paragraph, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    outline_level = OxmlElement("w:outlineLvl")
    outline_level.set(qn("w:val"), str(level))
    p_pr.append(outline_level)


def _write_deterministic_entry(
    archive: ZipFile, name: str, data: bytes, template: ZipInfo | None = None
) -> None:
    info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
    info.compress_type = ZIP_DEFLATED
    info.create_system = 3
    info.external_attr = (
        template.external_attr if template is not None else 0o600 << 16
    )
    archive.writestr(info, data)


def _normalize_package(path: Path, unsupported_part: bool) -> None:
    rewritten = path.with_suffix(".rewrite.docx")
    with ZipFile(path) as source, ZipFile(rewritten, "w") as target:
        for info in sorted(source.infolist(), key=lambda item: item.filename):
            _write_deterministic_entry(
                target, info.filename, source.read(info.filename), info
            )
        if unsupported_part:
            _write_deterministic_entry(
                target, "customXml/unsupported.bin", b"\x00opaque\xff"
            )
    rewritten.replace(path)


def build_docx(
    path: Path,
    paragraphs: list[dict[str, object]],
    *,
    add_table: bool = False,
    add_header_footer: bool = False,
    unsupported_part: bool = False,
) -> Path:
    document = Document()
    fixed_time = datetime(2024, 1, 1, tzinfo=UTC)
    document.core_properties.created = fixed_time
    document.core_properties.modified = fixed_time
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    for index, item in enumerate(paragraphs, start=1):
        paragraph = document.add_paragraph(str(item.get("text", "")))
        style = item.get("style")
        if style:
            if str(style) not in document.styles:
                document.styles.add_style(str(style), WD_STYLE_TYPE.PARAGRAPH)
            paragraph.style = str(style)
        fields = item.get("fields")
        if fields is None:
            field = item.get("field")
            fields = [field] if field else []
        for field in fields:
            _append_field(paragraph, str(field))
        bookmarks = item.get("bookmarks")
        if bookmarks is None:
            bookmark = item.get("bookmark")
            bookmarks = [bookmark] if bookmark else []
        for bookmark_offset, bookmark in enumerate(bookmarks):
            _append_bookmark(paragraph, str(bookmark), index * 100 + bookmark_offset)
        numbering_level = item.get("numbering_level")
        if numbering_level is not None:
            _set_numbering(paragraph, int(numbering_level))
        outline_level = item.get("outline_level")
        if outline_level is not None:
            _set_outline_level(paragraph, int(outline_level))
        if item.get("section_break"):
            document.add_section(WD_SECTION.NEW_PAGE)

    if add_table:
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
    if add_header_footer:
        section = document.sections[0]
        section.header.paragraphs[0].text = "Header text"
        _append_field(section.footer.paragraphs[0], " PAGE ")
        section.footer.paragraphs[0].add_run("Footer text")

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    _normalize_package(path, unsupported_part)
    return path


def build_minimal_roles_fixture(path: Path) -> Path:
    return build_docx(
        path,
        [
            {"text": "第一章 绪论", "style": "AppleHeading1"},
            {"text": "1.1 研究背景", "numbering_level": 1},
            {"text": "图 1-1 苹果树", "field": " SEQ Figure \\* ARABIC "},
            {"text": "可能是标题也可能是正文"},
            {"text": "参考文献"},
            {"text": "Apple Research Group. Orchard study."},
            {"text": "附录 A 调查表"},
        ],
        add_table=True,
        add_header_footer=True,
        unsupported_part=True,
    )
