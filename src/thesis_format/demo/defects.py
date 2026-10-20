"""Independent, one-concern corruptions for the deliberately broken demo."""

from __future__ import annotations

from collections import OrderedDict
from pathlib import Path
from shutil import copyfile
from typing import Callable

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


Defect = Callable[[str | Path, str | Path], Path]


def _open(source: str | Path, output: str | Path):
    source_path = Path(source)
    output_path = Path(output)
    if source_path.resolve() != output_path.resolve():
        output_path.parent.mkdir(parents=True, exist_ok=True)
        copyfile(source_path, output_path)
    return Document(output_path), output_path


def _save(document: Document, output: Path) -> Path:
    document.save(output)
    return output


def _style(document: Document, name: str):
    return document.styles[name]


def _ordered(parent, child) -> None:
    from thesis_format.demo.build import _insert_by_schema_order

    _insert_by_schema_order(parent, child)


def wrong_paper_size(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    for section in document.sections:
        section.page_width = Mm(215.9)
        section.page_height = Mm(279.4)
    return _save(document, destination)


def wrong_body_margins(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    for section in document.sections:
        section.left_margin = Mm(20)
        section.right_margin = Mm(20)
    return _save(document, destination)


def wrong_section_start(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    if len(document.sections) > 1:
        section_type = document.sections[1]._sectPr.find(qn("w:type"))
        if section_type is None:
            section_type = OxmlElement("w:type")
            _ordered(document.sections[1]._sectPr, section_type)
        section_type.set(qn("w:val"), "continuous")
    return _save(document, destination)


def mixed_fonts(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    style = _style(document, "AppleBody")
    fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        _ordered(style.element.get_or_add_rPr(), fonts)
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "SimSun")
    return _save(document, destination)


def manual_heading_numbers(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    for paragraph in document.paragraphs:
        if paragraph.style.name == "AppleHeading1":
            num_pr = paragraph._p.find(qn("w:pPr") + "/" + qn("w:numPr"))
            if num_pr is not None:
                paragraph._p.get_or_add_pPr().remove(num_pr)
            text = paragraph._p.find(".//" + qn("w:t"))
            if text is not None:
                text.text = "第一章 " + (text.text or "")
            break
    return _save(document, destination)


def unreliable_heading_styles(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    style = document.styles.add_style("UnreliableHeading", WD_STYLE_TYPE.PARAGRAPH)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "1")
    _ordered(style.element.get_or_add_pPr(), outline)
    for paragraph in document.paragraphs:
        if paragraph.style.name == "AppleHeading2":
            paragraph.style = style
            break
    return _save(document, destination)


def inconsistent_spacing(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    style = _style(document, "AppleBody")
    style.paragraph_format.space_after = Pt(9)
    style.paragraph_format.line_spacing = 1.5
    return _save(document, destination)


def missing_page_numbers(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            for run in list(paragraph._p):
                paragraph._p.remove(run)
    return _save(document, destination)


def broken_headers_footers(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    if len(document.sections) > 2:
        document.sections[2].header.is_linked_to_previous = False
        document.sections[2].header.paragraphs[0].text = "错误页眉"
        document.sections[2].footer.is_linked_to_previous = False
        document.sections[2].footer.paragraphs[0].text = "错误页脚"
    return _save(document, destination)


def malformed_toc(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    for paragraph in document.paragraphs:
        codes = paragraph._p.xpath(".//w:instrText")
        for code in codes:
            if code.text and 'TOC \\o "1-3"' in code.text:
                code.text = ' TOC \\o "1-2" '
                return _save(document, destination)
    return _save(document, destination)


def caption_positions(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    for paragraph in document.paragraphs[:-1]:
        if paragraph._p.xpath(".//w:drawing"):
            caption = paragraph._p.getnext()
            if caption is not None:
                paragraph._p.addprevious(caption)
            break
    return _save(document, destination)


def oversized_figure(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    extents = document._element.xpath(".//wp:extent")
    if extents:
        extents[0].set("cx", str(round(180 * 36000)))
    return _save(document, destination)


def table_border_noise(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    if document.tables:
        borders = document.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            _ordered(document.tables[0]._tbl.tblPr, borders)
        inside = OxmlElement("w:insideV")
        inside.set(qn("w:val"), "single")
        inside.set(qn("w:sz"), "4")
        borders.append(inside)
    return _save(document, destination)


def table_page_behavior(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    if document.tables:
        for row in document.tables[0].rows:
            props = row._tr.get_or_add_trPr()
            for child in list(props):
                if child.tag in {qn("w:tblHeader"), qn("w:cantSplit")}:
                    props.remove(child)
    return _save(document, destination)


def reference_indent(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    style = _style(document, "AppleReference")
    style.paragraph_format.first_line_indent = Mm(0)
    style.paragraph_format.left_indent = Mm(0)
    return _save(document, destination)


def direct_formatting(source: str | Path, output: str | Path) -> Path:
    document, destination = _open(source, output)
    for paragraph in document.paragraphs:
        if paragraph.style.name == "AppleHeading1" and paragraph.runs:
            paragraph.runs[0].italic = True
            break
    return _save(document, destination)


DEFECTS: "OrderedDict[str, Defect]" = OrderedDict(
    (
        ("wrong_paper_size", wrong_paper_size),
        ("wrong_body_margins", wrong_body_margins),
        ("wrong_section_start", wrong_section_start),
        ("mixed_fonts", mixed_fonts),
        ("manual_heading_numbers", manual_heading_numbers),
        ("unreliable_heading_styles", unreliable_heading_styles),
        ("inconsistent_spacing", inconsistent_spacing),
        ("missing_page_numbers", missing_page_numbers),
        ("broken_headers_footers", broken_headers_footers),
        ("malformed_toc", malformed_toc),
        ("caption_positions", caption_positions),
        ("oversized_figure", oversized_figure),
        ("table_border_noise", table_border_noise),
        ("table_page_behavior", table_page_behavior),
        ("reference_indent", reference_indent),
        ("direct_formatting", direct_formatting),
    )
)


__all__ = [*DEFECTS, "DEFECTS"]
