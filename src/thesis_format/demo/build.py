"""Build the anonymous Apple University template and demonstration thesis."""

from __future__ import annotations

import argparse
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from thesis_format.demo.content import Table, Thesis, build_demo_thesis
from thesis_format.demo.figures import generate_demo_figures
from thesis_format.profile import FormatProfile, Style, load_profile


ROOT = Path(__file__).resolve().parents[3]
DEFAULT_PROFILE = ROOT / "rules" / "apple-university.yaml"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

_PPR_ORDER = (
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
    "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
    "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
    "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
    "pPrChange",
)
_RPR_ORDER = (
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
    "strike", "dstrike", "outline", "shadow", "emboss", "imprint", "noProof",
    "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
    "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd",
    "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
    "specVanish", "oMath", "rPrChange",
)
_STYLE_ORDER = (
    "name", "aliases", "basedOn", "next", "link", "autoRedefine", "hidden",
    "uiPriority", "semiHidden", "unhideWhenUsed", "qFormat", "locked",
    "personal", "personalCompose", "personalReply", "rsid", "pPr", "rPr",
    "sectPr", "tblPr", "trPr", "tcPr",
)
_NUMBERING_ORDER = ("numPicBullet", "abstractNum", "numIdMacAtCleanup", "num")
_ABSTRACT_NUM_ORDER = (
    "nsid", "multiLevelType", "tmpl", "name", "styleLink", "numStyleLink", "lvl"
)
_LVL_ORDER = (
    "start", "numFmt", "lvlRestart", "pStyle", "isLgl", "suff", "lvlText",
    "lvlPicBulletId", "legacy", "lvlJc", "pPr", "rPr",
)
_SECTPR_ORDER = (
    "headerReference", "footerReference", "footnotePr", "endnotePr", "type",
    "pgSz", "pgMar", "paperSrc", "pgBorders", "lnNumType", "pgNumType",
    "cols", "formProt", "vAlign", "noEndnote", "titlePg", "textDirection",
    "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange",
)
_TCPR_ORDER = (
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
    "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange",
)
_TBLPR_ORDER = (
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
    "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
    "tblDescription", "tblPrChange",
)
_TRPR_ORDER = (
    "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
    "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
    "ins", "del", "trPrChange",
)
_ORDER_BY_PARENT = {
    qn("w:pPr"): _PPR_ORDER,
    qn("w:rPr"): _RPR_ORDER,
    qn("w:style"): _STYLE_ORDER,
    qn("w:numbering"): _NUMBERING_ORDER,
    qn("w:abstractNum"): _ABSTRACT_NUM_ORDER,
    qn("w:lvl"): _LVL_ORDER,
    qn("w:sectPr"): _SECTPR_ORDER,
    qn("w:tcPr"): _TCPR_ORDER,
    qn("w:tblPr"): _TBLPR_ORDER,
    qn("w:trPr"): _TRPR_ORDER,
}


def _local_name(element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def _insert_by_schema_order(parent, child) -> None:
    """Insert a child according to the ISO WordprocessingML sequence."""

    order = _ORDER_BY_PARENT.get(parent.tag)
    if order is None:
        parent.append(child)
        return
    try:
        child_rank = order.index(_local_name(child))
    except ValueError:
        parent.append(child)
        return
    for index, sibling in enumerate(parent):
        try:
            sibling_rank = order.index(_local_name(sibling))
        except ValueError:
            continue
        if sibling_rank > child_rank:
            parent.insert(index, child)
            return
    parent.append(child)


def _on(name: str, enabled: bool) -> OxmlElement:
    element = OxmlElement(name)
    if not enabled:
        element.set(qn("w:val"), "0")
    return element


def _field(paragraph, instruction: str, placeholder: str = "") -> None:
    def add_component(component) -> None:
        run = OxmlElement("w:r")
        run.append(component)
        paragraph._p.append(run)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    add_component(begin)

    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    add_component(code)

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    add_component(separate)

    text = OxmlElement("w:t")
    text.text = placeholder
    add_component(text)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    add_component(end)


def _bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def _replace_child(parent, child) -> None:
    existing = parent.find(child.tag)
    if existing is not None:
        parent.remove(existing)
    _insert_by_schema_order(parent, child)


def _configure_style(word_style, definition: Style) -> None:
    word_style.base_style = None
    p_pr = word_style.element.get_or_add_pPr()
    paragraph = definition.paragraph
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), str(round(paragraph.space_before_pt * 20)))
    spacing.set(qn("w:after"), str(round(paragraph.space_after_pt * 20)))
    if paragraph.line_spacing.type == "exact":
        spacing.set(qn("w:lineRule"), "exact")
        spacing.set(qn("w:line"), str(round(float(paragraph.line_spacing.value_pt) * 20)))
    else:
        spacing.set(qn("w:lineRule"), "auto")
        spacing.set(qn("w:line"), str(round(float(paragraph.line_spacing.value) * 240)))
    _replace_child(p_pr, spacing)

    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), str(round(paragraph.left_indent_mm * 1440 / 25.4)))
    if paragraph.first_line_chars:
        indent.set(qn("w:firstLineChars"), str(round(paragraph.first_line_chars * 100)))
    if paragraph.hanging_indent_mm:
        indent.set(
            qn("w:hanging"),
            str(round(paragraph.hanging_indent_mm * 1440 / 25.4)),
        )
    _replace_child(p_pr, indent)

    alignment = OxmlElement("w:jc")
    alignment.set(
        qn("w:val"),
        {"justify": "both", "left": "start", "right": "end"}.get(
            paragraph.alignment, paragraph.alignment
        ),
    )
    _replace_child(p_pr, alignment)
    _replace_child(p_pr, _on("w:keepNext", paragraph.keep_with_next))
    _replace_child(p_pr, _on("w:pageBreakBefore", paragraph.page_break_before))

    r_pr = word_style.element.get_or_add_rPr()
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), definition.run.font_latin)
    fonts.set(qn("w:hAnsi"), definition.run.font_latin)
    fonts.set(qn("w:eastAsia"), definition.run.font_zh)
    _replace_child(r_pr, fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(round(definition.run.size_pt * 2)))
    _replace_child(r_pr, size)
    _replace_child(r_pr, _on("w:b", definition.run.bold))
    _replace_child(r_pr, _on("w:i", definition.run.italic))


def _install_profile_styles(document: Document, profile: FormatProfile) -> None:
    for role, definition in profile.styles.items():
        style_name = definition.word_styles[0]
        if style_name in document.styles:
            style = document.styles[style_name]
        else:
            style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        _configure_style(style, definition)
        style.element.set(qn("w:customStyle"), "1")
        if role.startswith("heading_"):
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), str(int(role[-1]) - 1))
            _replace_child(style.element.get_or_add_pPr(), outline)


def _set_num_pr(paragraph, level: int, num_id: int = 10) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num))
    _replace_child(p_pr, num_pr)


def _install_multilevel_numbering(document: Document) -> None:
    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "10")
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    _insert_by_schema_order(abstract, multi)
    formats = (
        ("chineseCounting", "第%1章"),
        ("decimal", "%1.%2"),
        ("decimal", "%1.%2.%3"),
    )
    for level, (fmt, text) in enumerate(formats):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "space")
        style = OxmlElement("w:pStyle")
        style.set(qn("w:val"), f"AppleHeading{level + 1}")
        for child in (start, num_fmt, style, suffix, lvl_text):
            _insert_by_schema_order(lvl, child)
        _insert_by_schema_order(abstract, lvl)
    _insert_by_schema_order(numbering, abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "10")
    abstract_id = OxmlElement("w:abstractNumId")
    abstract_id.set(qn("w:val"), "10")
    num.append(abstract_id)
    _insert_by_schema_order(numbering, num)
    appendix_abstract = OxmlElement("w:abstractNum")
    appendix_abstract.set(qn("w:abstractNumId"), "11")
    appendix_level = OxmlElement("w:lvl")
    appendix_level.set(qn("w:ilvl"), "0")
    appendix_start = OxmlElement("w:start")
    appendix_start.set(qn("w:val"), "1")
    appendix_format = OxmlElement("w:numFmt")
    appendix_format.set(qn("w:val"), "upperLetter")
    appendix_text = OxmlElement("w:lvlText")
    appendix_text.set(qn("w:val"), "附录 %1")
    appendix_suffix = OxmlElement("w:suff")
    appendix_suffix.set(qn("w:val"), "space")
    appendix_style = OxmlElement("w:pStyle")
    appendix_style.set(qn("w:val"), "AppleAppendixHeading")
    for child in (
        appendix_start,
        appendix_format,
        appendix_style,
        appendix_suffix,
        appendix_text,
    ):
        _insert_by_schema_order(appendix_level, child)
    _insert_by_schema_order(appendix_abstract, appendix_level)
    _insert_by_schema_order(numbering, appendix_abstract)
    appendix_num = OxmlElement("w:num")
    appendix_num.set(qn("w:numId"), "11")
    appendix_id = OxmlElement("w:abstractNumId")
    appendix_id.set(qn("w:val"), "11")
    appendix_num.append(appendix_id)
    _insert_by_schema_order(numbering, appendix_num)


def _clean_metadata(document: Document) -> None:
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.comments = ""
    properties.category = ""
    properties.keywords = ""
    properties.subject = ""
    properties.title = ""
    properties.created = datetime(2000, 1, 1, tzinfo=UTC)
    properties.modified = datetime(2000, 1, 1, tzinfo=UTC)
    settings = document.settings.element
    zoom = settings.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    track = settings.find(qn("w:trackRevisions"))
    if track is not None:
        settings.remove(track)


def _set_page(section, profile: FormatProfile) -> None:
    section.page_width = Mm(profile.page.paper.width_mm)
    section.page_height = Mm(profile.page.paper.height_mm)
    section.top_margin = Mm(profile.page.body.top_mm)
    section.bottom_margin = Mm(profile.page.body.bottom_mm)
    section.left_margin = Mm(profile.page.body.left_mm)
    section.right_margin = Mm(profile.page.body.right_mm)
    section.header_distance = Mm(profile.page.header_distance_mm)
    section.footer_distance = Mm(profile.page.footer_distance_mm)


def _clear_story(story) -> None:
    for paragraph in story.paragraphs:
        paragraph._element.getparent().remove(paragraph._element)


def _header_footer(
    section,
    header_instruction: str | None,
    page_style: str | None,
    *,
    header_placeholder: str = "标题",
    header_number_instruction: str | None = None,
    header_number_placeholder: str = "",
) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    _clear_story(section.header)
    _clear_story(section.footer)
    if header_instruction:
        header = section.header.add_paragraph()
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if header_number_instruction:
            _field(header, header_number_instruction, header_number_placeholder)
            header.add_run(" ")
        _field(header, header_instruction, header_placeholder)
    if page_style:
        footer = section.footer.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(footer, "PAGE", "1")
        sect_pr = section._sectPr
        pg_num = sect_pr.find(qn("w:pgNumType"))
        if pg_num is None:
            pg_num = OxmlElement("w:pgNumType")
            _insert_by_schema_order(sect_pr, pg_num)
        pg_num.set(qn("w:fmt"), page_style)


def _restart_page_numbers(section, start: int = 1) -> None:
    pg_num = section._sectPr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        _insert_by_schema_order(section._sectPr, pg_num)
    pg_num.set(qn("w:start"), str(start))


def _new_document(profile: FormatProfile) -> Document:
    document = Document()
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    _clean_metadata(document)
    _install_profile_styles(document, profile)
    _install_multilevel_numbering(document)
    _set_page(document.sections[0], profile)
    return document


def _add_section(document: Document, profile: FormatProfile):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    _set_page(section, profile)
    return section


def _styled(document: Document, text: str, profile: FormatProfile, role: str):
    return document.add_paragraph(text, style=profile.styles[role].word_styles[0])


def _add_field_page(
    document: Document,
    profile: FormatProfile,
    title: str,
    instruction: str,
    bookmark_name: str | None = None,
    bookmark_id: int = 1,
) -> None:
    heading = _styled(document, title, profile, "front_heading")
    if bookmark_name:
        _bookmark(heading, bookmark_name, bookmark_id)
    field_paragraph = document.add_paragraph()
    _field(field_paragraph, instruction, "更新字段后显示")


def _caption_fields(paragraph, label: str) -> None:
    paragraph.add_run(f"{label} ")
    _field(paragraph, 'STYLEREF "AppleHeading1" \\n', "1")
    paragraph.add_run("-")
    _field(paragraph, f"SEQ {label} \\s 1 \\* ARABIC", "1")


def _caption(
    document: Document,
    title: str,
    profile: FormatProfile,
    role: str,
    label: str,
):
    paragraph = _styled(document, "", profile, role)
    _caption_fields(paragraph, label)
    paragraph.add_run(f" {title}")
    return paragraph


def _three_line_table(
    document: Document,
    table_model: Table,
    profile: FormatProfile,
    bookmark: tuple[str, int] | None = None,
):
    caption = _caption(
        document,
        table_model.caption,
        profile,
        "table_caption",
        "表",
    )
    if bookmark is not None:
        _bookmark(caption, *bookmark)
    table = document.add_table(rows=1, cols=len(table_model.headers))
    table.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    _insert_by_schema_order(table._tbl.tblPr, borders)
    header = table.rows[0]
    header_props = header._tr.get_or_add_trPr()
    _insert_by_schema_order(header_props, _on("w:tblHeader", True))
    _insert_by_schema_order(header_props, _on("w:cantSplit", True))
    for cell, value in zip(header.cells, table_model.headers, strict=True):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.style = profile.styles.table_body.word_styles[0]
        paragraph.add_run(value)
        tc_borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "000000")
        tc_borders.append(bottom)
        _insert_by_schema_order(cell._tc.get_or_add_tcPr(), tc_borders)
    for values in table_model.rows:
        row = table.add_row()
        _insert_by_schema_order(
            row._tr.get_or_add_trPr(), _on("w:cantSplit", True)
        )
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.style = profile.styles.table_body.word_styles[0]
            paragraph.add_run(value)
    if table_model.note:
        _styled(document, table_model.note, profile, "source_note")
    return table


def _set_alt_text(inline_shape, alt_text: str, title: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", title)


def _build_front(document: Document, profile: FormatProfile, thesis: Thesis | None) -> None:
    title_zh = thesis.title_zh if thesis else "中文论文题目"
    title_en = thesis.title_en if thesis else "English Thesis Title"
    cover = _styled(document, title_zh, profile, "title_zh")
    _bookmark(cover, "demo_cover", 1)
    _styled(document, title_en, profile, "title_en")
    _styled(document, "苹果大学（虚构）", profile, "front_heading")
    if thesis:
        _styled(document, "摘 要", profile, "front_heading")
        _styled(document, thesis.abstract_zh, profile, "abstract_zh")
        _styled(
            document,
            "关键词：" + "；".join(thesis.keywords_zh),
            profile,
            "keywords",
        )
        _styled(document, "ABSTRACT", profile, "front_heading")
        _styled(document, thesis.abstract_en, profile, "abstract_en")
        _styled(
            document,
            "Keywords: " + "; ".join(thesis.keywords_en),
            profile,
            "keywords",
        )
    else:
        _styled(document, "摘 要", profile, "front_heading")
        _styled(document, "在此填写中文摘要。", profile, "abstract_zh")
        _styled(document, "关键词：关键词一；关键词二", profile, "keywords")
        _styled(document, "ABSTRACT", profile, "front_heading")
        _styled(document, "Write the English abstract here.", profile, "abstract_en")
        _styled(document, "Keywords: keyword one; keyword two", profile, "keywords")
    _add_field_page(
        document,
        profile,
        "目录",
        'TOC \\o "1-3" \\h \\z \\u',
        "demo_toc",
        2,
    )
    _add_field_page(document, profile, "图目录", 'TOC \\h \\z \\c "图"')
    _add_field_page(document, profile, "表目录", 'TOC \\h \\z \\c "表"')
    _styled(document, "缩略语表", profile, "front_heading")
    _styled(document, "CV：变异系数", profile, "body")


def build_template(
    output_path: str | Path,
    profile: FormatProfile | None = None,
) -> Path:
    """Build an anonymous template entirely from profile data and code."""

    selected = profile or load_profile(DEFAULT_PROFILE)
    document = _new_document(selected)
    _build_front(document, selected, None)
    body_section = _add_section(document, selected)
    _header_footer(body_section, 'STYLEREF "AppleHeading1"', "decimal")
    _restart_page_numbers(body_section)
    body = _styled(document, "章节标题", selected, "heading_1")
    _set_num_pr(body, 0)
    _bookmark(body, "demo_body", 3)
    subsection = _styled(document, "二级标题", selected, "heading_2")
    _set_num_pr(subsection, 1)
    _styled(document, "在此填写正文。", selected, "body")
    figure = document.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.keep_with_next = True
    figure.add_run("[在此插入图形]")
    _bookmark(figure, "demo_figure", 4)
    _caption(document, "图题示例", selected, "figure_caption", "图")
    placeholder = Table(
        "表 1-1",
        ("列一", "列二"),
        (("内容", "内容"),),
        "表题示例",
        "",
    )
    _three_line_table(
        document,
        placeholder,
        selected,
        bookmark=("demo_table", 5),
    )
    references = _styled(document, "参考文献", selected, "front_heading")
    _bookmark(references, "demo_references", 6)
    _styled(document, "[1] 作者. 题名. 年份. 出版物.", selected, "reference")
    appendix_section = _add_section(document, selected)
    _header_footer(
        appendix_section,
        'STYLEREF "AppleAppendixHeading"',
        "decimal",
        header_placeholder="附录标题",
        header_number_instruction='STYLEREF "AppleAppendixHeading" \\n',
        header_number_placeholder="附录 A",
    )
    appendix = _styled(document, "附录标题", selected, "appendix_heading")
    _set_num_pr(appendix, 0, 11)
    _bookmark(appendix, "demo_appendix", 7)
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


def _chapter_content(
    document: Document,
    profile: FormatProfile,
    thesis: Thesis,
    asset_dir: Path,
) -> None:
    figures = iter(thesis.figures)
    tables = iter(thesis.tables)
    next_figure = next(figures, None)
    next_table = next(tables, None)
    font_candidates = (
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/System/Library/Fonts/Hiragino Sans GB.ttc"),
        Path("/System/Library/Fonts/STHeiti Medium.ttc"),
    )
    font_path = next((path for path in font_candidates if path.is_file()), None)
    generated = generate_demo_figures(asset_dir, font_path=font_path)
    figure_paths = dict(zip((item.number for item in thesis.figures), generated, strict=True))

    first_heading = True
    first_figure = True
    first_table = True
    for chapter in thesis.chapters:
        heading = _styled(document, chapter.title, profile, "heading_1")
        _set_num_pr(heading, 0)
        if first_heading:
            _bookmark(heading, "demo_body", 3)
            first_heading = False
        for paragraph_model in chapter.paragraphs:
            _styled(document, paragraph_model.text, profile, paragraph_model.role)
        for section_model in chapter.sections:
            section_heading = _styled(
                document, section_model.title, profile, "heading_2"
            )
            _set_num_pr(section_heading, 1)
            for paragraph_model in section_model.paragraphs:
                _styled(document, paragraph_model.text, profile, paragraph_model.role)
            if next_figure and next_figure.number.startswith(f"图 {chapter.number}-"):
                figure_paragraph = document.add_paragraph()
                figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                figure_paragraph.paragraph_format.keep_with_next = True
                shape = figure_paragraph.add_run().add_picture(
                    str(figure_paths[next_figure.number]), width=Mm(145)
                )
                _set_alt_text(shape, next_figure.alt_text, next_figure.number)
                if first_figure:
                    _bookmark(figure_paragraph, "demo_figure", 4)
                    first_figure = False
                _caption(
                    document,
                    next_figure.caption,
                    profile,
                    "figure_caption",
                    "图",
                )
                next_figure = next(figures, None)
            if next_table and next_table.number.startswith(f"表 {chapter.number}-"):
                _three_line_table(
                    document,
                    next_table,
                    profile,
                    bookmark=("demo_table", 5) if first_table else None,
                )
                first_table = False
                next_table = next(tables, None)


def build_thesis(
    output_path: str | Path,
    profile: FormatProfile | None = None,
    thesis: Thesis | None = None,
) -> Path:
    """Render the complete fictional Thesis model as a valid DOCX."""

    selected = profile or load_profile(DEFAULT_PROFILE)
    model = thesis or build_demo_thesis()
    document = _new_document(selected)
    asset_dir = Path(output_path).parent / "demo-assets"
    _build_front(document, selected, model)
    body_section = _add_section(document, selected)
    _header_footer(body_section, 'STYLEREF "AppleHeading1"', "decimal")
    _restart_page_numbers(body_section)
    _chapter_content(document, selected, model, asset_dir)

    references_section = _add_section(document, selected)
    _header_footer(references_section, 'STYLEREF "AppleFrontHeading"', "decimal")
    references = _styled(document, "参考文献", selected, "front_heading")
    _bookmark(references, "demo_references", 6)
    for reference in model.references:
        _styled(
            document,
            (
                f"[{reference.number}] {reference.author}. {reference.title}. "
                f"{reference.year}. {reference.publication}."
            ),
            selected,
            "reference",
        )

    appendix_section = _add_section(document, selected)
    _header_footer(
        appendix_section,
        'STYLEREF "AppleAppendixHeading"',
        "decimal",
        header_placeholder=model.appendix.title,
        header_number_instruction='STYLEREF "AppleAppendixHeading" \\n',
        header_number_placeholder="附录 A",
    )
    appendix = _styled(
        document,
        model.appendix.title,
        selected,
        "appendix_heading",
    )
    _set_num_pr(appendix, 0, 11)
    _bookmark(appendix, "demo_appendix", 7)
    for paragraph_model in model.appendix.paragraphs:
        _styled(document, paragraph_model.text, selected, paragraph_model.role)
    for table_model in model.appendix.tables:
        _three_line_table(document, table_model, selected)
    _styled(document, "致谢", selected, "front_heading")
    _styled(document, model.acknowledgements, selected, "body")
    _styled(document, model.data_disclosure, selected, "body")

    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


def build_before(
    output_path: str | Path,
    profile: FormatProfile | None = None,
) -> Path:
    """Build the readable but deliberately defective demonstration document."""

    selected = profile or load_profile(DEFAULT_PROFILE)
    destination = build_thesis(output_path, selected)
    from thesis_format.demo.defects import DEFECTS

    for defect in DEFECTS.values():
        defect(destination, destination)
    return destination


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--before", type=Path, required=True)
    parser.add_argument("--profile", type=Path, default=DEFAULT_PROFILE)
    return parser


def main(argv: Iterable[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    profile = load_profile(args.profile)
    build_template(args.template, profile)
    build_before(args.before, profile)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
